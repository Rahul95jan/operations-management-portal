import json
import io
import re
import csv
import pandas as pd
from datetime import datetime, timedelta
from dateutil import parser as date_parser
from fastapi.responses import FileResponse, Response

from fastapi import FastAPI, Depends, Form, File, UploadFile
from typing import Optional
from pydantic import BaseModel
from fastapi.middleware.cors import CORSMiddleware
from models.operations import OperationsAnalytics

from models.resource import Resource
from models.resource_requirement import ResourceRequirement
from models.audit_log import AuditLog
from resource_schemas import (
    ResourceUpdate,
    ResourceRequirementCreate,
    ResourceRequirementStatusUpdate,
    AppSettingsUpdate,
)
from app_settings import get_settings, settings_to_dict
from services.storage import save_file, resolve_path
from resource_tracking import (
    build_tracking_table,
    refresh_requirement_status,
    compute_delay,
    session_resource_detail,
    list_pending_requirements,
)
from resource_downloads import build_lms_package, lms_filename
from models.resource_download_log import ResourceDownloadLog
from models.resource_email_log import ResourceEmailLog
from resource_emails import initial_request_email, reminder_email, confirmation_email, ops_notification_email
from email_service import send_email as send_generic_email
from scheduler import start_scheduler, stop_scheduler, run_reminder_check, REMINDER_INTERVAL_HOURS
from resource_analytics import (
    compute_resource_analytics,
    compute_mentor_performance,
    compute_mentor_heatmap,
    compute_at_risk_mentors,
)
from resource_export import build_export_rows
from resource_tokens import get_or_create_session_token, get_session_by_token, submission_url
from resource_scheduling import compute_due_at
from resource_defaults import DEFAULT_REQUIREMENTS
from mentor_performance import get_mentor_scorecard, get_mentor_trend, DIMENSION_WEIGHTS, CLASSIFICATION_BANDS
import session_reports
from schemas import (
    SessionReportUpdate,
    ReportStatusUpdate,
    SessionAttendanceCreate,
    SessionAttendanceUpdate,
)
import os

OPS_NOTIFICATION_EMAIL = os.getenv("OPS_NOTIFICATION_EMAIL")

from database import SessionLocal
from database import get_db as get_auth_db  # same session-factory auth.py's get_current_user uses,
# so a user fetched via get_current_user and a db handed to the same endpoint are the same
# SQLAlchemy session — required for FastAPI to dedupe the dependency and avoid mixing sessions
# (main.py's own get_db below is a separate, functionally-identical function used everywhere
# else in this file; treating it as interchangeable with database.get_db caused
# "Instance is not persistent within this Session" errors on profile updates/photo uploads).
from sqlalchemy.orm import Session
from sqlalchemy import func

from models.user import User
from models.session import Session as SessionModel
from models.mentor import Mentor
from models.batch import Batch
from models.zoom_account import ZoomAccount
from models.invoice import Invoice
from models.nps import NPSFeedback
from schemas import NPSCreate
from models.session_analytics import SessionAnalytics
from database import engine
from models.user import Base
from pdf_generator import generate_invoice, generate_webinar_report, generate_nps_report, generate_analytics_report, generate_mentor_performance_report, generate_webinar_report_pdf, generate_session_report_pdf, generate_session_reports_list_pdf
from nps_insights import compute_nps_insights
from models.invoice import Invoice
from schemas import InvoiceCreate, InvoicePaymentUpdate
import os

from email_service import send_invoice_email
from models.invoice import Invoice
print("Invoice Columns:", Invoice.__table__.columns.keys())
from email_service import send_invoice_email



from schemas import (
    UserCreate,
    UserUpdate,
    LoginRequest,
    AdminUserCreate,
    AdminUserPermissionsUpdate,
    ZoomAccountCreate,
    ZoomAccountUpdate,
    SessionCreate,
    MentorCreate,
    BatchCreate,
    InvoiceCreate,
    SessionAnalyticsCreate,
    ZoomAnalyticsCreate,
)
from auth import (
    hash_password,
    verify_password,
    create_access_token,
    get_current_user,
    require_permission,
    require_super_admin,
    user_public,
    has_permission,
    count_active_super_admins,
    SECTIONS,
    ASSIGNABLE_SECTIONS,
    SUPER_ADMIN,
    ADMIN,
)

app = FastAPI()


@app.on_event("startup")
def _start_resource_scheduler():
    start_scheduler()


@app.on_event("shutdown")
def _stop_resource_scheduler():
    stop_scheduler()


def get_db():
    db = SessionLocal()

    try:
        yield db

    finally:
        db.close()

# ==========================
# CORS
# ==========================

allowed_origins_env = os.getenv("ALLOWED_ORIGINS")
if allowed_origins_env:
    allowed_origins = [o.strip() for o in allowed_origins_env.split(",") if o.strip()]
else:
    allowed_origins = [
        "http://localhost:3000",
        "http://127.0.0.1:3000",
    ]

app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# An unhandled exception that reaches Starlette's default error handling can end up
# bypassing CORSMiddleware entirely, so the browser sees a response with no
# Access-Control-Allow-Origin header and reports a confusing "blocked by CORS policy" /
# "Failed to fetch" — hiding the real 500 underneath it. This catch-all makes sure any
# unhandled exception still comes back as a normal, CORS-safe JSON error instead.
@app.exception_handler(Exception)
async def _unhandled_exception_handler(request, exc: Exception):
    import traceback
    from fastapi.responses import JSONResponse

    print(f"❌ Unhandled exception on {request.method} {request.url.path}: {exc}")
    traceback.print_exc()
    return JSONResponse(
        status_code=500,
        content={"detail": "Something went wrong on the server. Please try again."},
    )

# ==========================
# HOME
# ==========================

@app.get("/")
def home():
    return {
        "message": "Operations Portal API Running"
    }

# ==========================
# USERS
# ==========================

@app.post("/users")
def create_user(user: UserCreate):
    return {
        "message": "User Created",
        "data": user
    }


@app.post("/auth/login")
def login(payload: LoginRequest, db: Session = Depends(get_db)):
    from fastapi import HTTPException

    user = db.query(User).filter(User.username == payload.username).first()
    if not user or not verify_password(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Invalid username or password.")
    if not user.is_active:
        raise HTTPException(status_code=403, detail="This account has been deactivated.")

    token = create_access_token(user)
    return {"token": token, "user": user_public(user)}


@app.get("/auth/me")
def read_auth_me(user: User = Depends(get_current_user)):
    return user_public(user)


@app.get("/auth/permission-catalog")
def permission_catalog(user: User = Depends(get_current_user)):
    return {"sections": SECTIONS, "assignable_sections": ASSIGNABLE_SECTIONS}


@app.get("/users/me")
def get_my_profile(user: User = Depends(get_current_user)):
    return user_public(user)


@app.put("/users/me")
def update_my_profile(payload: UserUpdate, user: User = Depends(get_current_user), db: Session = Depends(get_auth_db)):
    if payload.name is not None:
        user.name = payload.name
    if payload.email is not None:
        user.email = payload.email
    if payload.phone is not None:
        user.phone = payload.phone
    # Role is never self-editable — only a Super Admin changes roles, via User Management.
    db.commit()
    db.refresh(user)
    return user_public(user)


@app.post("/users/me/photo")
async def upload_my_photo(file: UploadFile = File(...), user: User = Depends(get_current_user), db: Session = Depends(get_auth_db)):
    from fastapi import HTTPException

    if not file.content_type or not file.content_type.startswith("image/"):
        raise HTTPException(status_code=400, detail="Only image files (JPG, PNG, WEBP) are allowed.")

    contents = await file.read()
    if len(contents) > 2 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Image must be 2MB or smaller.")
    file.file.seek(0)

    file_meta = save_file(file)
    user.photo_path = file_meta["file_path"]
    db.commit()

    return {"message": "Photo uploaded successfully", "has_photo": True}


@app.get("/users/me/photo")
def get_my_photo(user: User = Depends(get_current_user)):
    from fastapi import HTTPException

    resolved = resolve_path(user.photo_path) if user.photo_path else None
    if not resolved or not os.path.exists(resolved):
        raise HTTPException(status_code=404, detail="No photo found for this account.")

    return FileResponse(resolved)


@app.get("/users/{user_id}/photo")
def get_user_photo(user_id: int, db: Session = Depends(get_auth_db)):
    # Deliberately unauthenticated, same as /mentors/{id}/photo — a plain <img src> can never
    # send a Bearer token (browsers don't support custom headers on image loads), so this is
    # the URL every <img> tag in the app actually points at; POST /users/me/photo (uploading)
    # still requires being logged in as that account.
    from fastapi import HTTPException

    user = db.query(User).filter(User.id == user_id).first()
    resolved = resolve_path(user.photo_path) if user and user.photo_path else None
    if not user or not resolved or not os.path.exists(resolved):
        raise HTTPException(status_code=404, detail="No photo found for this user.")

    return FileResponse(resolved)


# ==========================
# ADMINISTRATION — user management & permissions (Super Admin only)
# ==========================

@app.get("/admin/users")
def list_users(current: User = Depends(require_super_admin), db: Session = Depends(get_db)):
    users = db.query(User).order_by(User.id).all()
    return [user_public(u) for u in users]


@app.post("/admin/users")
def create_admin_user(payload: AdminUserCreate, current: User = Depends(require_super_admin), db: Session = Depends(get_db)):
    from fastapi import HTTPException

    if payload.role not in (SUPER_ADMIN, ADMIN):
        raise HTTPException(status_code=400, detail="Role must be ADMIN or SUPER_ADMIN.")
    if db.query(User).filter(User.username == payload.username).first():
        raise HTTPException(status_code=400, detail="That username is already taken.")
    if db.query(User).filter(User.email == payload.email).first():
        raise HTTPException(status_code=400, detail="That email is already registered.")

    valid_perms = {f"{s}.{a}" for s, meta in ASSIGNABLE_SECTIONS.items() for a in meta["actions"]}
    perms = [] if payload.role == SUPER_ADMIN else [p for p in payload.permissions if p in valid_perms]

    new_user = User(
        name=payload.name,
        email=payload.email,
        username=payload.username,
        password_hash=hash_password(payload.password),
        role=payload.role,
        permissions=perms,
        is_active=True,
    )
    db.add(new_user)
    db.commit()
    db.refresh(new_user)

    _log_audit(db, "User Created", details=f"Role: {payload.role}, Permissions: {len(perms)} section action(s)", user=current.name)

    return user_public(new_user)


@app.put("/admin/users/{user_id}/access")
def update_user_access(user_id: int, payload: AdminUserPermissionsUpdate, current: User = Depends(require_super_admin), db: Session = Depends(get_db)):
    from fastapi import HTTPException

    target = db.query(User).filter(User.id == user_id).first()
    if not target:
        raise HTTPException(status_code=404, detail="User not found.")
    if target.id == current.id:
        raise HTTPException(status_code=400, detail="You can't change your own role or permissions here.")

    demoting = payload.role is not None and target.role == SUPER_ADMIN and payload.role != SUPER_ADMIN
    deactivating = payload.is_active is False and target.is_active and target.role == SUPER_ADMIN
    if (demoting or deactivating) and count_active_super_admins(db, exclude_id=target.id) == 0:
        raise HTTPException(status_code=400, detail="At least one active Super Admin must remain.")

    changes = []
    if payload.role is not None and payload.role != target.role:
        if payload.role not in (SUPER_ADMIN, ADMIN):
            raise HTTPException(status_code=400, detail="Role must be ADMIN or SUPER_ADMIN.")
        changes.append(f"Role: {target.role} → {payload.role}")
        target.role = payload.role
        if payload.role == SUPER_ADMIN:
            target.permissions = []

    if payload.permissions is not None and target.role != SUPER_ADMIN:
        valid_perms = {f"{s}.{a}" for s, meta in ASSIGNABLE_SECTIONS.items() for a in meta["actions"]}
        new_perms = sorted(set(p for p in payload.permissions if p in valid_perms))
        old_perms = set(target.permissions or [])
        granted = sorted(set(new_perms) - old_perms)
        revoked = sorted(old_perms - set(new_perms))
        if granted:
            changes.append(f"Granted: {', '.join(granted)}")
        if revoked:
            changes.append(f"Revoked: {', '.join(revoked)}")
        target.permissions = new_perms

    if payload.is_active is not None and payload.is_active != target.is_active:
        changes.append(f"Status: {'Active' if payload.is_active else 'Deactivated'}")
        target.is_active = payload.is_active

    db.commit()
    db.refresh(target)

    if changes:
        _log_audit(db, "Permission Updated", details=f"Target: {target.name} — {'; '.join(changes)}", user=current.name)

    return user_public(target)


@app.get("/admin/activity-logs")
def get_activity_logs(current: User = Depends(require_super_admin), db: Session = Depends(get_db)):
    logs = db.query(AuditLog).order_by(AuditLog.created_at.desc()).limit(200).all()
    return [
        {
            "id": l.id,
            "timestamp": l.created_at.isoformat() if l.created_at else None,
            "performed_by": l.user,
            "action": l.action,
            "details": l.details,
        }
        for l in logs
    ]


# ==========================
# ZOOM ACCOUNTS — the pool of Zoom logins/profiles sessions get run under
# ==========================

@app.get("/zoom-accounts")
def get_zoom_accounts(db: Session = Depends(get_db), _user: User = Depends(require_permission("sessions", "view"))):
    accounts = db.query(ZoomAccount).filter(ZoomAccount.is_active == True).order_by(ZoomAccount.id).all()  # noqa: E712
    return accounts


@app.post("/zoom-accounts")
def create_zoom_account(payload: ZoomAccountCreate, db: Session = Depends(get_db), _user: User = Depends(require_permission("sessions", "create"))):
    from fastapi import HTTPException

    email = payload.email.strip()
    if not email:
        raise HTTPException(status_code=400, detail="Zoom ID can't be empty.")

    existing = db.query(ZoomAccount).filter(func.lower(ZoomAccount.email) == email.lower()).first()
    if existing:
        if not existing.is_active:
            existing.is_active = True
            db.commit()
            db.refresh(existing)
        return existing

    account = ZoomAccount(email=email)
    db.add(account)
    db.commit()
    db.refresh(account)
    return account


@app.put("/zoom-accounts/{account_id}")
def update_zoom_account(account_id: int, payload: ZoomAccountUpdate, db: Session = Depends(get_db), _user: User = Depends(require_permission("sessions", "edit"))):
    from fastapi import HTTPException

    email = payload.email.strip()
    if not email:
        raise HTTPException(status_code=400, detail="Zoom ID can't be empty.")

    account = db.query(ZoomAccount).filter(ZoomAccount.id == account_id).first()
    if not account:
        raise HTTPException(status_code=404, detail="Zoom ID not found.")

    duplicate = (
        db.query(ZoomAccount)
        .filter(func.lower(ZoomAccount.email) == email.lower(), ZoomAccount.id != account_id)
        .first()
    )
    if duplicate:
        raise HTTPException(status_code=400, detail="That Zoom ID already exists.")

    account.email = email
    db.commit()
    db.refresh(account)
    return account


@app.delete("/zoom-accounts/{account_id}")
def delete_zoom_account(account_id: int, db: Session = Depends(get_db), _user: User = Depends(require_permission("sessions", "edit"))):
    account = db.query(ZoomAccount).filter(ZoomAccount.id == account_id).first()
    if account:
        # Soft-delete — existing sessions may still reference this email by value, so we
        # keep the row but drop it from the pick-list instead of hard-deleting.
        account.is_active = False
        db.commit()
    return {"message": "Zoom ID removed"}


# ==========================
# SESSIONS
# ==========================

def _notify_mentor_session_change(db, session_obj, action, old_date=None, old_time=None):
    """Emails the mentor (address pulled fresh from Mentor Management, never a stale
    stored value) whenever a session they're assigned to is scheduled, rescheduled, or
    cancelled/deleted. Reuses the same SMTP config as invoice emails — see email_service.py.
    Never raises: a notification failure must never block the session CRUD it's attached to."""
    from email_service import send_session_notification

    if not session_obj.mentor_name:
        return {"notified": False, "reason": "no_mentor_assigned"}

    mentor = db.query(Mentor).filter(Mentor.name == session_obj.mentor_name).first()
    mentor_email = mentor.email if mentor else None
    if not mentor_email:
        print(f"⚠️  No email on file for mentor '{session_obj.mentor_name}' — skipping session notification.")
        return {"notified": False, "reason": "no_mentor_email"}

    id_line = f"Zoom ID: {session_obj.zoom_id}" if session_obj.zoom_id else (f"Webinar ID: {session_obj.webinar_id}" if session_obj.webinar_id else None)

    if action == "Scheduled":
        subject = f"Session Scheduled: {session_obj.topic}"
        body = (
            f"Hi {session_obj.mentor_name},\n\n"
            f"A session has been scheduled for you.\n\n"
            f"Topic: {session_obj.topic}\n"
            f"Batch: {session_obj.batch_name or 'N/A (Webinar)'}\n"
            f"Date: {session_obj.session_date}\n"
            f"Time: {session_obj.session_time}\n"
            + (f"{id_line}\n" if id_line else "")
            + f"\nThis session is now blocked on the Operations Portal calendar.\n\n"
            f"Regards,\nKrish Naik Academy Team"
        )
    elif action == "Rescheduled":
        subject = f"Session Rescheduled: {session_obj.topic}"
        body = (
            f"Hi {session_obj.mentor_name},\n\n"
            f"Your session has been rescheduled.\n\n"
            f"Topic: {session_obj.topic}\n"
            f"Batch: {session_obj.batch_name or 'N/A (Webinar)'}\n"
            f"Previous: {old_date} {old_time}\n"
            f"New: {session_obj.session_date} {session_obj.session_time}\n"
            + (f"{id_line}\n" if id_line else "")
            + f"\nPlease update your calendar accordingly.\n\n"
            f"Regards,\nKrish Naik Academy Team"
        )
    else:  # Cancelled
        subject = f"Session Cancelled: {session_obj.topic}"
        body = (
            f"Hi {session_obj.mentor_name},\n\n"
            f"The following session has been cancelled and removed from the schedule.\n\n"
            f"Topic: {session_obj.topic}\n"
            f"Batch: {session_obj.batch_name or 'N/A (Webinar)'}\n"
            f"Was scheduled for: {session_obj.session_date} {session_obj.session_time}\n\n"
            f"Regards,\nKrish Naik Academy Team"
        )

    sent = send_session_notification(mentor_email, subject, body)
    return {"notified": sent, "reason": "sent" if sent else "email_unavailable", "mentor_email": mentor_email}


@app.post("/sessions")
def create_session(session: SessionCreate, _user: User = Depends(require_permission("sessions", "create"))):

    db = SessionLocal()

    new_session = SessionModel(
        project_name=session.topic,
        topic=session.topic,
        mentor_name=session.mentor_name,
        batch_name=session.batch_name,
        session_date=session.session_date,
        session_time=session.session_time,
        status=session.status,
        session_type=session.session_type,
        webinar_id=session.webinar_id,
        zoom_id=session.zoom_id,
        remarks=session.remarks
    )

    db.add(new_session)
    db.commit()
    db.refresh(new_session)

    if new_session.status == "Completed":
        _ensure_default_requirements(db, new_session)

    notification = None
    if new_session.status == "Scheduled":
        notification = _notify_mentor_session_change(db, new_session, "Scheduled")

    db.close()

    return {
        "message": "Session Created Successfully",
        "id": new_session.id,
        "calendar_block": {
            "blocked": new_session.status == "Scheduled",
            "mentor_name": new_session.mentor_name,
            "date": new_session.session_date,
            "time": new_session.session_time,
            "notification": notification,
        },
    }


@app.get("/sessions")
def get_sessions(db: Session = Depends(get_db), _user: User = Depends(require_permission("sessions", "view"))):

    sessions = db.query(SessionModel).all()

    return sessions

@app.put("/sessions/{session_id}")
def update_session(session_id: int, session: SessionCreate, _user: User = Depends(require_permission("sessions", "edit"))):

    db = SessionLocal()

    existing_session = (
        db.query(SessionModel)
        .filter(SessionModel.id == session_id)
        .first()
    )

    if not existing_session:
        db.close()
        return {"message": "Session Not Found"}

    previous_status = existing_session.status
    old_date, old_time, old_status = existing_session.session_date, existing_session.session_time, existing_session.status
    just_cancelled = session.status == "Cancelled" and old_status != "Cancelled"
    rescheduled = (session.session_date != old_date or session.session_time != old_time) and old_status != "Cancelled" and session.status != "Cancelled"

    existing_session.topic = session.topic
    existing_session.mentor_name = session.mentor_name
    existing_session.batch_name = session.batch_name
    existing_session.session_date = session.session_date
    existing_session.session_time = session.session_time
    existing_session.status = session.status
    existing_session.session_type = session.session_type
    existing_session.webinar_id = session.webinar_id
    existing_session.zoom_id = session.zoom_id
    if session.remarks is not None:
        existing_session.remarks = session.remarks

    db.commit()
    db.refresh(existing_session)

    if existing_session.status == "Completed" and previous_status != "Completed":
        _ensure_default_requirements(db, existing_session)

    notification = None
    if rescheduled:
        notification = _notify_mentor_session_change(db, existing_session, "Rescheduled", old_date=old_date, old_time=old_time)
    elif just_cancelled:
        notification = _notify_mentor_session_change(db, existing_session, "Cancelled")

    db.close()

    return {
        "message": "Session Updated Successfully",
        "calendar_block": {
            "blocked": existing_session.status == "Scheduled",
            "mentor_name": existing_session.mentor_name,
            "date": existing_session.session_date,
            "time": existing_session.session_time,
            "notification": notification,
        },
    }


@app.delete("/sessions/{session_id}")
def delete_session(session_id: int, _user: User = Depends(require_permission("sessions", "delete"))):
    from fastapi import HTTPException
    from sqlalchemy import or_
    from sqlalchemy.exc import IntegrityError

    db = SessionLocal()

    session = (
        db.query(SessionModel)
        .filter(SessionModel.id == session_id)
        .first()
    )

    if session:
        _notify_mentor_session_change(db, session, "Cancelled")

        # Resources/requirements/their email + download logs are foreign-keyed to
        # sessions.id (unlike session_reports/session_attendance/session_analytics/
        # zoom_analytics, which are deliberately plain int columns so that historical
        # reports/analytics survive a session being deleted). Those FK'd rows can't
        # outlive their session, so they're cleaned up first, children before parents,
        # or Postgres rejects the session delete with a ForeignKeyViolation.
        requirement_ids = [
            r.id for r in db.query(ResourceRequirement.id).filter(ResourceRequirement.session_id == session_id).all()
        ]
        resource_ids = [
            r.id for r in db.query(Resource.id).filter(Resource.session_id == session_id).all()
        ]

        db.query(ResourceEmailLog).filter(
            or_(
                ResourceEmailLog.session_id == session_id,
                ResourceEmailLog.resource_requirement_id.in_(requirement_ids),
                ResourceEmailLog.resource_id.in_(resource_ids),
            )
        ).delete(synchronize_session=False)

        db.query(ResourceDownloadLog).filter(
            or_(
                ResourceDownloadLog.session_id == session_id,
                ResourceDownloadLog.resource_id.in_(resource_ids),
            )
        ).delete(synchronize_session=False)

        db.query(Resource).filter(Resource.session_id == session_id).delete(synchronize_session=False)
        db.query(ResourceRequirement).filter(ResourceRequirement.session_id == session_id).delete(synchronize_session=False)

        db.delete(session)
        try:
            db.commit()
        except IntegrityError as e:
            db.rollback()
            db.close()
            raise HTTPException(
                status_code=409,
                detail="Can't delete this session — other records still reference it. Please try again or contact support.",
            ) from e

    db.close()

    return {
        "message": "Session Deleted Successfully"
    }

# ==========================
# MENTORS
# ==========================

@app.post("/mentors")
def create_mentor(mentor: MentorCreate, _user: User = Depends(require_permission("mentors", "create"))):

    db = SessionLocal()

    new_mentor = Mentor(
        name=mentor.name,
        email=mentor.email,
        phone=mentor.phone,
        expertise=mentor.expertise,
        linkedin=mentor.linkedin,
        hourly_rate=mentor.hourly_rate,
        status=mentor.status or "Active"
    )

    db.add(new_mentor)
    db.commit()
    db.refresh(new_mentor)

    db.close()

    return {
        "message": "Mentor Created Successfully",
        "id": new_mentor.id
    }


@app.get("/mentors")
def get_mentors(db: Session = Depends(get_db), _user: User = Depends(require_permission("mentors", "view"))):

    mentors = db.query(Mentor).all()

    return mentors


@app.put("/mentors/{mentor_id}")
def update_mentor(mentor_id: int, mentor: MentorCreate, _user: User = Depends(require_permission("mentors", "edit"))):

    db = SessionLocal()

    existing_mentor = (
        db.query(Mentor)
        .filter(Mentor.id == mentor_id)
        .first()
    )

    if not existing_mentor:
        db.close()
        return {
            "message": "Mentor Not Found"
        }

    existing_mentor.name = mentor.name
    existing_mentor.email = mentor.email
    existing_mentor.phone = mentor.phone
    existing_mentor.expertise = mentor.expertise
    existing_mentor.linkedin = mentor.linkedin
    existing_mentor.hourly_rate = mentor.hourly_rate
    existing_mentor.status = mentor.status or existing_mentor.status or "Active"

    db.commit()
    db.refresh(existing_mentor)

    db.close()

    return {
        "message": "Mentor Updated Successfully"
    }


@app.delete("/mentors/{mentor_id}")
def delete_mentor(mentor_id: int, _user: User = Depends(require_permission("mentors", "delete"))):

    db = SessionLocal()

    mentor = (
        db.query(Mentor)
        .filter(Mentor.id == mentor_id)
        .first()
    )

    if mentor:
        db.delete(mentor)
        db.commit()

    db.close()

    return {
        "message": "Mentor Deleted Successfully"
    }


@app.post("/mentors/{mentor_id}/photo")
async def upload_mentor_photo(mentor_id: int, file: UploadFile = File(...), db: Session = Depends(get_db)):
    from fastapi import HTTPException

    mentor = db.query(Mentor).filter(Mentor.id == mentor_id).first()
    if not mentor:
        raise HTTPException(status_code=404, detail="Mentor not found.")

    if not file.content_type or not file.content_type.startswith("image/"):
        raise HTTPException(status_code=400, detail="Only image files (JPG, PNG, WEBP) are allowed.")

    contents = await file.read()
    if len(contents) > 2 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Image must be 2MB or smaller.")
    file.file.seek(0)

    file_meta = save_file(file)
    mentor.photo_path = file_meta["file_path"]
    db.commit()

    return {"message": "Photo uploaded successfully", "has_photo": True}


@app.get("/mentors/{mentor_id}/photo")
def get_mentor_photo(mentor_id: int, db: Session = Depends(get_db)):
    from fastapi import HTTPException

    mentor = db.query(Mentor).filter(Mentor.id == mentor_id).first()
    resolved = resolve_path(mentor.photo_path) if mentor and mentor.photo_path else None
    if not mentor or not resolved or not os.path.exists(resolved):
        raise HTTPException(status_code=404, detail="No photo found for this mentor.")

    return FileResponse(resolved)

# ==========================
# BATCHES
# ==========================

@app.post("/batches")
def create_batch(batch: BatchCreate, _user: User = Depends(require_permission("batches", "create"))):

    db = SessionLocal()

    new_batch = Batch(
        batch_name=batch.batch_name,
        course_name=batch.course_name,
        strength=batch.strength,
        mentor_name=batch.mentor_name,
        status=batch.status or "Active"
    )

    db.add(new_batch)
    db.commit()
    db.refresh(new_batch)

    db.close()

    return {
        "message": "Batch Created Successfully",
        "id": new_batch.id
    }
@app.put("/batches/{batch_id}")
def update_batch(batch_id: int, batch: BatchCreate, _user: User = Depends(require_permission("batches", "edit"))):

    db = SessionLocal()

    existing_batch = (
        db.query(Batch)
        .filter(Batch.id == batch_id)
        .first()
    )

    if not existing_batch:
        db.close()
        return {
            "message": "Batch Not Found"
        }

    existing_batch.batch_name = batch.batch_name
    existing_batch.course_name = batch.course_name
    existing_batch.strength = batch.strength
    existing_batch.mentor_name = batch.mentor_name
    existing_batch.status = batch.status or existing_batch.status or "Active"

    db.commit()
    db.refresh(existing_batch)

    db.close()

    return {
        "message": "Batch Updated Successfully"
    }

@app.get("/batches")
def get_batches(db: Session = Depends(get_db), _user: User = Depends(require_permission("batches", "view"))):

    batches = db.query(Batch).all()

    return batches

@app.delete("/batches/{batch_id}")
def delete_batch(batch_id: int, _user: User = Depends(require_permission("batches", "delete"))):

    db = SessionLocal()

    batch = (
        db.query(Batch)
        .filter(Batch.id == batch_id)
        .first()
    )

    if batch:
        db.delete(batch)
        db.commit()

    db.close()

    return {
        "message": "Batch Deleted Successfully"
    }

# ==========================
# DASHBOARD
# ==========================

@app.get("/dashboard")
def dashboard(db: Session = Depends(get_db)):

    data = {
        "total_users": db.query(User).count(),
        "total_sessions": db.query(SessionModel).count(),
        "total_mentors": db.query(Mentor).count(),
        "total_batches": db.query(Batch).count()
    }

    return data

# ==========================
# ANALYTICS
# ==========================

@app.get("/recent-sessions")
def recent_sessions():

    db = SessionLocal()

    sessions = (
        db.query(SessionModel)
        .order_by(SessionModel.id.desc())
        .limit(5)
        .all()
    )

    db.close()

    return sessions
@app.get("/recent-batches")
def recent_batches():

    db = SessionLocal()

    batches = (
        db.query(Batch)
        .order_by(Batch.id.desc())
        .limit(5)
        .all()
    )

    db.close()

    return batches

@app.get("/analytics")
def analytics():

    db = SessionLocal()

    data = {
        "total_users": db.query(User).count(),
        "total_sessions": db.query(SessionModel).count(),
        "total_mentors": db.query(Mentor).count(),
        "total_batches": db.query(Batch).count(),

        "scheduled_sessions":
            db.query(SessionModel)
            .filter(SessionModel.status == "Scheduled")
            .count(),

        "completed_sessions":
            db.query(SessionModel)
            .filter(SessionModel.status == "Completed")
            .count()
    }

    db.close()

    return data

@app.get("/recent-mentors")
def recent_mentors():

    db = SessionLocal()

    mentors = (
        db.query(Mentor)
        .order_by(Mentor.id.desc())
        .limit(5)
        .all()
    )

    db.close()

    return mentors

@app.get("/upcoming-sessions")
def upcoming_sessions():

    db = SessionLocal()

    sessions = (
        db.query(SessionModel)
        .filter(SessionModel.status == "Scheduled")
        .order_by(SessionModel.session_date.asc())
        .limit(5)
        .all()
    )

    db.close()

    return sessions

@app.get("/export-sessions")
def export_sessions():

    db = SessionLocal()

    sessions = db.query(SessionModel).all()

    data = []

    for session in sessions:
        data.append({
            "ID": session.id,
            "Topic": session.topic,
            "Mentor": session.mentor_name,
            "Batch": session.batch_name,
            "Date": session.session_date,
            "Time": session.session_time,
            "Status": session.status
        })

    df = pd.DataFrame(data)

    file_name = "sessions.xlsx"

    df.to_excel(
        file_name,
        index=False
    )

    db.close()

    return FileResponse(
        file_name,
        filename=file_name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.get("/export-batches")
def export_batches():

    db = SessionLocal()

    batches = db.query(Batch).all()

    data = []

    for batch in batches:
        data.append({
            "ID": batch.id,
            "Batch": batch.batch_name,
            "Course": batch.course_name,
            "Strength": batch.strength,
            "Mentor": batch.mentor_name,
            "Status": batch.status
        })

    df = pd.DataFrame(data)

    file_name = "batches.xlsx"

    df.to_excel(
        file_name,
        index=False
    )

    db.close()

    return FileResponse(
        file_name,
        filename=file_name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.get("/export-invoices")
def export_invoices():

    db = SessionLocal()

    invoices = db.query(Invoice).all()

    data = []

    for inv in invoices:
        data.append({
            "Invoice No.": inv.invoice_number,
            "Mentor": inv.mentor_name,
            "Batch": inv.batch_name,
            "Month": inv.month,
            "Sessions": inv.total_sessions,
            "Hours": inv.total_hours,
            "Hourly Rate": inv.hourly_rate,
            "Amount": inv.total_amount,
            "Invoice Date": inv.created_at.strftime("%Y-%m-%d") if inv.created_at else None,
            "Due Date": inv.due_date,
            "Status": inv.payment_status,
            "Payment Date": inv.payment_date,
            "Payment Mode": inv.payment_mode,
        })

    df = pd.DataFrame(data)

    file_name = "invoices.xlsx"

    df.to_excel(
        file_name,
        index=False
    )

    db.close()

    return FileResponse(
        file_name,
        filename=file_name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.get("/dashboard/overview")
def dashboard_overview(db: Session = Depends(get_db)):
    print("1")

    total_mentors = db.query(Mentor).count()
    print("2")

    total_batches = db.query(Batch).count()
    print("3")

    total_sessions = db.query(SessionModel).count()
    print("4")

    scheduled_sessions = (
        db.query(SessionModel)
        .filter(SessionModel.status == "Scheduled")
        .count()
    )
    print("5")

    completed_sessions = (
        db.query(SessionModel)
        .filter(SessionModel.status == "Completed")
        .count()
    )
    print("6")

    active_mentors = (
        db.query(Mentor)
        .filter(Mentor.status == "Active")
        .count()
    )
    print("7")

    return {
        "total_mentors": total_mentors,
        "total_batches": total_batches,
        "total_sessions": total_sessions,
        "scheduled_sessions": scheduled_sessions,
        "completed_sessions": completed_sessions,
        "active_mentors": active_mentors,
    }
@app.get("/dashboard/session-analytics")
def session_analytics(
    course_name: Optional[str] = None,
    mentor_name: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    db: Session = Depends(get_db),
):
    base_query = db.query(SessionModel)
    if course_name:
        base_query = base_query.filter(SessionModel.course_name == course_name)
    if mentor_name:
        base_query = base_query.filter(SessionModel.mentor_name == mentor_name)
    if date_from:
        base_query = base_query.filter(SessionModel.session_date >= date_from)
    if date_to:
        base_query = base_query.filter(SessionModel.session_date <= date_to)

    total_sessions = base_query.count()
    completed_sessions = base_query.filter(SessionModel.status == "Completed").count()
    scheduled_sessions = base_query.filter(SessionModel.status == "Scheduled").count()
    cancelled_sessions = base_query.filter(SessionModel.status == "Cancelled").count()

    return {
        "total_sessions": total_sessions,
        "completed_sessions": completed_sessions,
        "scheduled_sessions": scheduled_sessions,
        "cancelled_sessions": cancelled_sessions
    }
@app.get("/dashboard/mentor-analytics")
def mentor_analytics(
    course_name: Optional[str] = None,
    mentor_name: Optional[str] = None,
    db: Session = Depends(get_db),
):
    query = db.query(Mentor)

    if course_name:
        course_mentor_names = [
            b.mentor_name for b in db.query(Batch).filter(Batch.course_name == course_name).all()
            if b.mentor_name
        ]
        query = query.filter(Mentor.name.in_(course_mentor_names))

    if mentor_name:
        query = query.filter(Mentor.name == mentor_name)

    total_mentors = query.count()
    active_mentors = query.filter(Mentor.status == "Active").count()
    inactive_mentors = query.filter(Mentor.status == "Inactive").count()
    mentors = query.all()

    total_rate = 0

    for mentor in mentors:
        try:
            total_rate += int(mentor.hourly_rate)
        except:
            pass

    average_hourly_rate = (
        total_rate / total_mentors
        if total_mentors > 0
        else 0
    )

    return {
        "total_mentors": total_mentors,
        "active_mentors": active_mentors,
        "inactive_mentors": inactive_mentors,
        "average_hourly_rate": round(average_hourly_rate, 2)
    }
@app.get("/dashboard/batch-analytics")
def batch_analytics(
    course_name: Optional[str] = None,
    mentor_name: Optional[str] = None,
    db: Session = Depends(get_db),
):
    query = db.query(Batch)
    if course_name:
        query = query.filter(Batch.course_name == course_name)
    if mentor_name:
        query = query.filter(Batch.mentor_name == mentor_name)

    total_batches = query.count()

    batches = query.all()

    total_strength = 0

    for batch in batches:
        if batch.strength:
            total_strength += batch.strength

    average_batch_strength = (
        total_strength / total_batches
        if total_batches > 0
        else 0
    )

    return {
        "total_batches": total_batches,
        "total_strength": total_strength,
        "average_batch_strength": round(
            average_batch_strength,
            2
        )
    }
@app.get("/dashboard/revenue-analytics")
def revenue_analytics(db: Session = Depends(get_db)):

    invoices = db.query(Invoice).all()

    total_invoices = len(invoices)

    total_revenue = 0
    paid_amount = 0
    pending_amount = 0

    for invoice in invoices:

        try:
            amount = int(invoice.total_amount)
        except:
            amount = 0

        total_revenue += amount

        if invoice.payment_status == "Paid":
            paid_amount += amount
        else:
            pending_amount += amount

    return {
        "total_invoices": total_invoices,
        "total_revenue": total_revenue,
        "paid_amount": paid_amount,
        "pending_amount": pending_amount
    }
@app.get("/dashboard/top-mentors")

def top_mentors(
    course_name: Optional[str] = None,
    mentor_name: Optional[str] = None,
    db: Session = Depends(get_db),
):
    mentor_query = db.query(Mentor)

    if course_name:
        course_mentor_names = [
            b.mentor_name for b in db.query(Batch).filter(Batch.course_name == course_name).all()
            if b.mentor_name
        ]
        mentor_query = mentor_query.filter(Mentor.name.in_(course_mentor_names))

    if mentor_name:
        mentor_query = mentor_query.filter(Mentor.name == mentor_name)

    mentors = mentor_query.all()

    leaderboard = []

    for mentor in mentors:

        session_query = db.query(SessionModel).filter(SessionModel.mentor_name == mentor.name)
        if course_name:
            session_query = session_query.filter(SessionModel.course_name == course_name)

        session_count = session_query.count()

        try:
            rate = int(mentor.hourly_rate)
        except:
            rate = 0

        revenue = session_count * 2 * rate

        leaderboard.append({
            "mentor_name": mentor.name,
            "sessions": session_count,
            "hourly_rate": rate,
            "revenue": revenue
        })

    leaderboard = sorted(
        leaderboard,
        key=lambda x: x["revenue"],
        reverse=True
    )

    return leaderboard


@app.get("/dashboard/top-batches")
def top_batches(
    course_name: Optional[str] = None,
    mentor_name: Optional[str] = None,
    db: Session = Depends(get_db),
):
    query = db.query(Batch)
    if course_name:
        query = query.filter(Batch.course_name == course_name)
    if mentor_name:
        query = query.filter(Batch.mentor_name == mentor_name)

    batches = query.all()

    leaderboard = []

    for batch in batches:

        session_count = (
            db.query(SessionModel)
            .filter(SessionModel.batch_name == batch.batch_name)
            .count()
        )

        leaderboard.append({
            "batch_name": batch.batch_name,
            "course_name": batch.course_name,
            "mentor_name": batch.mentor_name,
            "strength": batch.strength,
            "sessions": session_count
        })

    leaderboard = sorted(
        leaderboard,
        key=lambda x: x["sessions"],
        reverse=True
    )

    return leaderboard


@app.post("/invoices")
def create_invoice(invoice: InvoiceCreate):

    db = SessionLocal()

    try:
        # =====================================
        # Get the next invoice number
        # =====================================

        last_invoice = (
            db.query(Invoice)
            .order_by(Invoice.id.desc())
            .first()
        )

        next_number = (last_invoice.id + 1) if last_invoice else 1
        invoice_number = f"INV-{next_number:05d}"

        # =====================================
        # Create Invoice
        # =====================================

        due_date = invoice.due_date
        if not due_date:
            due_date = (datetime.utcnow() + timedelta(days=15)).strftime("%Y-%m-%d")

        new_invoice = Invoice(
            invoice_number=invoice_number,
            mentor_name=invoice.mentor_name,
            mentor_email=invoice.mentor_email,
            batch_name=invoice.batch_name,
            month=invoice.month,
            total_sessions=invoice.total_sessions,
            total_hours=invoice.total_hours,
            hourly_rate=invoice.hourly_rate,
            total_amount=invoice.total_amount,
            payment_status="Pending",
            due_date=due_date,
            notes=invoice.notes,
        )

        db.add(new_invoice)
        db.commit()
        db.refresh(new_invoice)

        # =====================================
        # Debug
        # =====================================

        print("=" * 60)
        print("Invoice ID      :", new_invoice.id)
        print("Invoice Number  :", new_invoice.invoice_number)
        print("Mentor          :", new_invoice.mentor_name)
        print("Payment Status  :", new_invoice.payment_status)
        print("=" * 60)

        # =====================================
        # Generate PDF
        # =====================================

        pdf_path = generate_invoice(new_invoice)

        print("PDF Generated:", pdf_path)

        return {
            "message": "Invoice Created Successfully",
            "invoice_id": new_invoice.id,
            "invoice_number": new_invoice.invoice_number,
            "payment_status": new_invoice.payment_status,
            "pdf": pdf_path,
        }

    except Exception as e:
        db.rollback()
        print("Invoice Error:", e)
        return {"error": str(e)}

    finally:
        db.close()


# =====================================================
# Download Invoice (Always Generate Latest PDF)
# =====================================================

@app.get("/download-invoice/{invoice_id}")
def download_invoice(invoice_id: int):

    db = SessionLocal()

    try:
        # Fetch latest invoice from database
        invoice = (
            db.query(Invoice)
            .filter(Invoice.id == invoice_id)
            .first()
        )

        if not invoice:
            return {"error": "Invoice not found"}

        print("=" * 60)
        print("Downloading Invoice")
        print("Invoice Number :", invoice.invoice_number)
        print("Payment Status :", invoice.payment_status)
        print("=" * 60)

        # Always regenerate PDF using latest database values
        pdf_path = generate_invoice(invoice)

        return FileResponse(
            path=pdf_path,
            filename=f"{invoice.invoice_number}.pdf",
            media_type="application/pdf",
        )

    except Exception as e:
        print("Download Error:", e)
        return {"error": str(e)}

    finally:
        db.close()


# ---------------------------------------------------------
# Send Invoice Email
# ---------------------------------------------------------

from fastapi import HTTPException
import traceback
import os

@app.post("/send-invoice/{invoice_id}")
def send_invoice(invoice_id: int):

    db = SessionLocal()

    try:
        # Get Invoice
        invoice = (
            db.query(Invoice)
            .filter(Invoice.id == invoice_id)
            .first()
        )

        if not invoice:
            return {
                "success": False,
                "message": "Invoice not found"
            }

        # Get Mentor
        mentor = (
            db.query(Mentor)
            .filter(Mentor.name == invoice.mentor_name)
            .first()
        )

        if not mentor:
            return {
                "success": False,
                "message": "Mentor not found"
            }

        # Get email
        mentor_email = invoice.mentor_email or mentor.email

        if not mentor_email:
            return {
                "success": False,
                "message": "Mentor email not found"
            }

        # Save email into invoice if missing
        if invoice.mentor_email != mentor_email:
            invoice.mentor_email = mentor_email
            db.commit()

        pdf_path = f"pdfs/invoice_{invoice_id}.pdf"

        if not os.path.exists(pdf_path):
            return {
                "success": False,
                "message": f"PDF not found: {pdf_path}"
            }

        print("=" * 60)
        print("Invoice ID :", invoice.id)
        print("Mentor     :", invoice.mentor_name)
        print("Email      :", mentor_email)
        print("PDF        :", pdf_path)
        print("=" * 60)

        success = send_invoice_email(
            receiver_email=mentor_email,
            pdf_path=pdf_path,
            invoice_number=invoice.invoice_number,
        )

        print("send_invoice_email returned:", success)

        if success:
            return {
                "success": True,
                "message": "Invoice sent successfully"
            }

        return {
            "success": False,
            "message": "send_invoice_email() returned False"
        }

    except Exception as e:
        traceback.print_exc()

        return {
            "success": False,
            "message": str(e)
        }

    finally:
        db.close()

@app.get("/invoices")
def get_invoices(db: Session = Depends(get_db)):

    invoices = db.query(Invoice).all()

    result = []

    for invoice in invoices:
        result.append({
            "id": invoice.id,
            "invoice_number": invoice.invoice_number,
            "mentor_name": invoice.mentor_name,
            "mentor_email": invoice.mentor_email,
            "batch_name": invoice.batch_name,
            "month": invoice.month,
            "total_sessions": invoice.total_sessions,
            "total_hours": invoice.total_hours,
            "hourly_rate": invoice.hourly_rate,
            "total_amount": invoice.total_amount,
            "payment_status": invoice.payment_status,
            "invoice_date": invoice.created_at.strftime("%Y-%m-%d") if invoice.created_at else None,
            "due_date": invoice.due_date,
            "payment_date": invoice.payment_date,
            "payment_mode": invoice.payment_mode,
            "transaction_id": invoice.transaction_id,
            "payment_reference": invoice.payment_reference,
            "notes": invoice.notes,
        })

    return result

@app.get("/batch-names")
def get_batch_names():

    db = SessionLocal()

    try:
        batches = db.query(Batch).all()
        return batches
    finally:
        db.close()

@app.get("/invoice-summary")
def invoice_summary():
    db = SessionLocal()

    invoices = db.query(Invoice).all()

    total_amount = 0
    total_hours = 0
    total_sessions = 0
    pending_amount = 0
    total_invoices = len(invoices)

    for invoice in invoices:

        total_amount += float(invoice.total_amount or 0)
        total_hours += float(invoice.total_hours or 0)
        total_sessions += invoice.total_sessions or 0

        if invoice.payment_status == "Pending":
            pending_amount += float(invoice.total_amount or 0)

    db.close()

    return {
        "total_amount": total_amount,
        "pending_amount": pending_amount,
        "total_hours": total_hours,
        "total_sessions": total_sessions,
        "total_invoices": total_invoices
    }

@app.put("/invoice-paid/{invoice_id}")
def mark_invoice_paid(invoice_id: int, payment: InvoicePaymentUpdate | None = None):

    db = SessionLocal()

    invoice = db.query(Invoice).filter(
        Invoice.id == invoice_id
    ).first()

    if not invoice:
        db.close()
        return {"message": "Invoice not found"}

    invoice.payment_status = "Paid"
    invoice.payment_date = (payment.payment_date if payment and payment.payment_date else None) or datetime.utcnow().strftime("%Y-%m-%d")
    if payment and payment.payment_mode:
        invoice.payment_mode = payment.payment_mode
    if payment and payment.transaction_id:
        invoice.transaction_id = payment.transaction_id
    if payment and payment.payment_reference:
        invoice.payment_reference = payment.payment_reference

    db.commit()
    db.refresh(invoice)

    db.close()

    return {
        "message": "Invoice marked as Paid"
    }
    
@app.get("/invoice/{invoice_id}")
def get_invoice(invoice_id: int):

    db = SessionLocal()

    invoice = db.query(Invoice).filter(
        Invoice.id == invoice_id
    ).first()

    db.close()

    return invoice

@app.put("/invoice/{invoice_id}")
def update_invoice(invoice_id: int, updated_invoice: InvoiceCreate):

    db = SessionLocal()

    invoice = db.query(Invoice).filter(
        Invoice.id == invoice_id
    ).first()

    if not invoice:
        db.close()
        return {"message": "Invoice not found"}

    invoice.mentor_name = updated_invoice.mentor_name
    invoice.batch_name = updated_invoice.batch_name
    invoice.month = updated_invoice.month
    invoice.total_sessions = updated_invoice.total_sessions
    invoice.total_hours = updated_invoice.total_hours
    invoice.hourly_rate = updated_invoice.hourly_rate
    invoice.total_amount = updated_invoice.total_amount
    invoice.payment_status = updated_invoice.payment_status
    if updated_invoice.due_date is not None:
        invoice.due_date = updated_invoice.due_date
    if updated_invoice.notes is not None:
        invoice.notes = updated_invoice.notes

    db.commit()
    db.refresh(invoice)

    db.close()

    return {
        "message": "Invoice Updated Successfully"
    }

@app.delete("/invoices/{invoice_id}")
def delete_invoice(invoice_id: int):

    db = SessionLocal()

    invoice = db.query(Invoice).filter(Invoice.id == invoice_id).first()

    if not invoice:
        db.close()
        return {"message": "Invoice not found"}

    db.delete(invoice)
    db.commit()
    db.close()

    return {"message": "Invoice deleted successfully"}

# ==========================================================
# NPS FEEDBACK
# ==========================================================

@app.post("/nps")
def create_nps(data: NPSCreate):

    db = SessionLocal()

    try:

        existing = (
            db.query(NPSFeedback)
            .filter(
                (NPSFeedback.learner_email == data.learner_email) |
                (NPSFeedback.mobile_number == data.mobile_number)
            )
            .first()
        )

        if existing:
            return {
                "success": False,
                "message": "Email or Mobile Number already exists."
            }

        feedback = NPSFeedback(
            learner_name=data.learner_name,
            learner_email=data.learner_email,
            mobile_number=data.mobile_number,
            course_name=data.course_name,
            batch_name=data.batch_name,
            mentor_name=data.mentor_name,
            instructor_rating=data.instructor_rating,
            doubt_rating=data.doubt_rating,
            website_rating=data.website_rating,
            nps_score=data.nps_score,
            feedback=data.feedback,
        )

        db.add(feedback)
        db.commit()
        db.refresh(feedback)

        return {
            "success": True,
            "message": "NPS Feedback Submitted Successfully",
            "id": feedback.id
        }

    finally:
        db.close()


def _filtered_nps_query(db, course_name=None, batch_name=None, mentor_name=None):
    query = db.query(NPSFeedback)

    if course_name:
        query = query.filter(NPSFeedback.course_name == course_name)
    if batch_name:
        query = query.filter(NPSFeedback.batch_name == batch_name)
    if mentor_name:
        query = query.filter(NPSFeedback.mentor_name == mentor_name)

    return query


@app.get("/nps")
def get_nps(course_name: str = None, batch_name: str = None, mentor_name: str = None):

    db = SessionLocal()

    try:

        feedbacks = _filtered_nps_query(db, course_name, batch_name, mentor_name).all()

        return [
            {
                "id": item.id,
                "learner_name": item.learner_name,
                "learner_email": item.learner_email,
                "mobile_number": item.mobile_number,
                "course_name": item.course_name,
                "batch_name": item.batch_name,
                "mentor_name": item.mentor_name,
                "instructor_rating": item.instructor_rating,
                "doubt_rating": item.doubt_rating,
                "website_rating": item.website_rating,
                "nps_score": item.nps_score,
                "feedback": item.feedback,
                "created_at": item.created_at,
            }
            for item in feedbacks
        ]

    finally:
        db.close()


@app.get("/nps/{feedback_id}")
def get_nps_by_id(feedback_id: int):

    db = SessionLocal()

    try:

        feedback = (
            db.query(NPSFeedback)
            .filter(NPSFeedback.id == feedback_id)
            .first()
        )

        if not feedback:
            return {"message": "Feedback Not Found"}

        return feedback

    finally:
        db.close()


@app.put("/nps/{feedback_id}")
def update_nps(feedback_id: int, data: NPSCreate):

    db = SessionLocal()

    try:
        feedback = (
            db.query(NPSFeedback)
            .filter(NPSFeedback.id == feedback_id)
            .first()
        )

        if not feedback:
            return {"message": "Feedback Not Found"}

        # Check if email or mobile already exists for another record
        existing = (
            db.query(NPSFeedback)
            .filter(
                (
                    (NPSFeedback.learner_email == data.learner_email) |
                    (NPSFeedback.mobile_number == data.mobile_number)
                ) &
                (NPSFeedback.id != feedback_id)
            )
            .first()
        )

        if existing:
            return {
                "message": "Email or Mobile Number already exists."
            }

        feedback.learner_name = data.learner_name
        feedback.learner_email = data.learner_email
        feedback.mobile_number = data.mobile_number
        feedback.course_name = data.course_name
        feedback.batch_name = data.batch_name
        feedback.mentor_name = data.mentor_name
        feedback.instructor_rating = data.instructor_rating
        feedback.doubt_rating = data.doubt_rating
        feedback.website_rating = data.website_rating
        feedback.nps_score = data.nps_score
        feedback.feedback = data.feedback

        db.commit()
        db.refresh(feedback)

        return {
            "message": "Feedback Updated Successfully"
        }

    finally:
        db.close()

@app.delete("/nps/{feedback_id}")
def delete_nps(feedback_id: int):

    db = SessionLocal()

    try:

        feedback = (
            db.query(NPSFeedback)
            .filter(NPSFeedback.id == feedback_id)
            .first()
        )

        if not feedback:
            return {"message": "Feedback Not Found"}

        db.delete(feedback)
        db.commit()

        return {
            "message": "Feedback Deleted Successfully"
        }

    finally:
        db.close()


# ==========================================================
# NPS Insights & Export
# ==========================================================

@app.get("/nps-insights")
def nps_insights(course_name: str = None, batch_name: str = None, mentor_name: str = None):

    db = SessionLocal()

    try:
        records = _filtered_nps_query(db, course_name, batch_name, mentor_name).all()
        return compute_nps_insights(records)

    finally:
        db.close()


@app.get("/export-nps")
def export_nps(course_name: str = None, batch_name: str = None, mentor_name: str = None):

    db = SessionLocal()

    try:
        records = _filtered_nps_query(db, course_name, batch_name, mentor_name).all()

        data = []

        for r in records:
            data.append({
                "ID": r.id,
                "Learner Name": r.learner_name,
                "Email": r.learner_email,
                "Mobile": r.mobile_number,
                "Course": r.course_name,
                "Batch": r.batch_name,
                "Mentor": r.mentor_name,
                "Instructor Rating": r.instructor_rating,
                "Doubt Rating": r.doubt_rating,
                "Website Rating": r.website_rating,
                "NPS Score": r.nps_score,
                "Feedback": r.feedback,
                "Submitted": r.created_at,
            })

        df = pd.DataFrame(data)

        file_name = "nps_responses.xlsx"
        df.to_excel(file_name, index=False)

        return FileResponse(
            file_name,
            filename=file_name,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    finally:
        db.close()


@app.get("/export-nps-report")
def export_nps_report(course_name: str = None, batch_name: str = None, mentor_name: str = None):

    db = SessionLocal()

    try:
        records = _filtered_nps_query(db, course_name, batch_name, mentor_name).all()
        insights = compute_nps_insights(records)

        pdf_path = generate_nps_report(records, insights)

        return FileResponse(
            path=pdf_path,
            filename="NPS_Analytics_Report.pdf",
            media_type="application/pdf",
        )

    finally:
        db.close()

@app.post("/session-analytics")
def create_session_analytics(data: SessionAnalyticsCreate):

    db = SessionLocal()

    analytics = SessionAnalytics(
        session_id=data.session_id,
        mentor_name=data.mentor_name,
        batch_name=data.batch_name,
        scheduled_duration=data.scheduled_duration,
        actual_duration=data.actual_duration,
        mentor_join_time=data.mentor_join_time,
        mentor_leave_time=data.mentor_leave_time,
        first_student_join=data.first_student_join,
        last_student_leave=data.last_student_leave,
        registered_count=data.registered_count,
        joined_count=data.joined_count,
        unique_attendance=data.unique_attendance,
        duplicate_joins=data.duplicate_joins,
        attendance_percentage=data.attendance_percentage,
        late_joiners=data.late_joiners,
        early_leavers=data.early_leavers,
        poll_count=data.poll_count,
        poll_participants=data.poll_participants,
        quiz_count=data.quiz_count,
        quiz_participants=data.quiz_participants,
        average_quiz_score=data.average_quiz_score,
        chat_messages=data.chat_messages,
        private_messages=data.private_messages,
        public_messages=data.public_messages,
        questions_asked=data.questions_asked,
        questions_answered=data.questions_answered,
        recording_available=data.recording_available,
        recording_duration=data.recording_duration,
    )

    db.add(analytics)
    db.commit()
    db.refresh(analytics)
    db.close()

    return {
        "message": "Session Analytics Created Successfully",
        "id": analytics.id,
    }

@app.get("/session-analytics")
def get_session_analytics():

    db = SessionLocal()

    analytics = db.query(SessionAnalytics).all()

    db.close()

    return analytics

@app.put("/session-analytics/{analytics_id}")
def update_session_analytics(
    analytics_id: int,
    data: SessionAnalyticsCreate
):

    db = SessionLocal()

    analytics = (
        db.query(SessionAnalytics)
        .filter(SessionAnalytics.id == analytics_id)
        .first()
    )

    if not analytics:
        db.close()
        return {"message": "Analytics Not Found"}

    analytics.session_id = data.session_id
    analytics.mentor_name = data.mentor_name
    analytics.batch_name = data.batch_name
    analytics.scheduled_duration = data.scheduled_duration
    analytics.actual_duration = data.actual_duration
    analytics.mentor_join_time = data.mentor_join_time
    analytics.mentor_leave_time = data.mentor_leave_time
    analytics.first_student_join = data.first_student_join
    analytics.last_student_leave = data.last_student_leave
    analytics.registered_count = data.registered_count
    analytics.joined_count = data.joined_count
    analytics.unique_attendance = data.unique_attendance
    analytics.duplicate_joins = data.duplicate_joins
    analytics.attendance_percentage = data.attendance_percentage
    analytics.late_joiners = data.late_joiners
    analytics.early_leavers = data.early_leavers
    analytics.poll_count = data.poll_count
    analytics.poll_participants = data.poll_participants
    analytics.quiz_count = data.quiz_count
    analytics.quiz_participants = data.quiz_participants
    analytics.average_quiz_score = data.average_quiz_score
    analytics.chat_messages = data.chat_messages
    analytics.private_messages = data.private_messages
    analytics.public_messages = data.public_messages
    analytics.questions_asked = data.questions_asked
    analytics.questions_answered = data.questions_answered
    analytics.recording_available = data.recording_available
    analytics.recording_duration = data.recording_duration

    db.commit()
    db.refresh(analytics)
    db.close()

    return {
        "message": "Session Analytics Updated Successfully"
    }   


@app.delete("/session-analytics/{analytics_id}")
def delete_session_analytics(analytics_id: int):

    db = SessionLocal()

    analytics = (
        db.query(SessionAnalytics)
        .filter(SessionAnalytics.id == analytics_id)
        .first()
    )

    if not analytics:
        db.close()
        return {"message": "Session Analytics Not Found"}

    db.delete(analytics)
    db.commit()
    db.close()

    return {
        "message": "Session Analytics Deleted Successfully"
    }
@app.get("/session-analytics-summary")
def session_analytics_summary():

    db = SessionLocal()

    analytics = db.query(SessionAnalytics).all()

    total_sessions = len(analytics)

    registered = sum(a.registered_count or 0 for a in analytics)

    joined = sum(a.joined_count or 0 for a in analytics)

    unique_attendance = sum(a.unique_attendance or 0 for a in analytics)

    mentor_hours = sum(a.mentor_total_hours or 0 for a in analytics)

    student_hours = sum(a.students_total_hours or 0 for a in analytics)

    total_duration = sum(a.actual_duration or 0 for a in analytics)

    attendance_percentage = (
        round((joined / registered) * 100, 2)
        if registered > 0
        else 0
    )

    db.close()

    return {
        "total_sessions": total_sessions,
        "registered": registered,
        "joined": joined,
        "unique_attendance": unique_attendance,
        "attendance_percentage": attendance_percentage,
        "mentor_hours": mentor_hours,
        "student_hours": student_hours,
        "total_duration": total_duration,
    }

@app.get("/attendance-trend")
def attendance_trend():

    db = SessionLocal()

    analytics = db.query(SessionAnalytics).all()

    chart = []

    for item in analytics:
        chart.append({
            "session": item.session_id,
            "attendance": item.attendance_percentage
        })

    db.close()

    return chart

@app.get("/poll-summary")
def poll_summary():

    db = SessionLocal()

    analytics = db.query(SessionAnalytics).all()

    total_polls = sum(a.poll_count or 0 for a in analytics)

    total_participants = sum(a.poll_participants or 0 for a in analytics)

    average_participation = (
        round(total_participants / total_polls, 2)
        if total_polls > 0
        else 0
    )

    average_accuracy = (
        round(
            sum(a.poll_accuracy or 0 for a in analytics) / len(analytics),
            2
        )
        if analytics
        else 0
    )

    db.close()

    return {
        "total_polls": total_polls,
        "total_participants": total_participants,
        "average_participation": average_participation,
        "average_accuracy": average_accuracy,
    }
@app.get("/quiz-summary")
def quiz_summary():

    db = SessionLocal()

    analytics = db.query(SessionAnalytics).all()

    total_quizzes = sum(a.quiz_count or 0 for a in analytics)

    total_participants = sum(a.quiz_participants or 0 for a in analytics)

    average_score = (
        round(
            sum(a.average_quiz_score or 0 for a in analytics) /
            len(analytics),
            2
        )
        if analytics
        else 0
    )

    db.close()

    return {
        "total_quizzes": total_quizzes,
        "total_participants": total_participants,
        "average_score": average_score,
    }
@app.get("/mentor-performance")
def mentor_performance():

    db = SessionLocal()

    analytics = db.query(SessionAnalytics).all()

    mentor_data = {}

    for a in analytics:

        mentor = a.mentor_name

        if mentor not in mentor_data:
            mentor_data[mentor] = {
                "mentor_name": mentor,
                "sessions": 0,
                "attendance": 0,
                "polls": 0,
                "hours": 0,
            }

        mentor_data[mentor]["sessions"] += 1
        mentor_data[mentor]["attendance"] += a.attendance_percentage or 0
        mentor_data[mentor]["polls"] += a.poll_count or 0
        mentor_data[mentor]["hours"] += a.mentor_total_hours or 0

    result = []

    for mentor in mentor_data.values():

        mentor["attendance"] = round(
            mentor["attendance"] / mentor["sessions"], 2
        )

        result.append(mentor)

    db.close()

    return result

@app.get("/poll-trend")
def poll_trend():

    db = SessionLocal()

    analytics = db.query(SessionAnalytics).all()

    result = []

    for item in analytics:
        result.append({
            "session": item.session_id,
            "polls": item.poll_count
        })

    db.close()

    return result


@app.get("/batch-analytics")
def batch_analytics(course_name: Optional[str] = None, mentor_name: Optional[str] = None):

    db = SessionLocal()

    query = db.query(Batch)
    if course_name:
        query = query.filter(Batch.course_name == course_name)
    if mentor_name:
        query = query.filter(Batch.mentor_name == mentor_name)

    batches = query.all()

    total_batches = len(batches)

    completed_batches = len(
        [b for b in batches if b.status == "Completed"]
    )

    ongoing_batches = len(
        [b for b in batches if b.status == "Ongoing"]
    )

    delayed_batches = len(
        [b for b in batches if b.status == "Delayed"]
    )

    average_attendance = (
        round(
            sum(b.attendance_percentage or 0 for b in batches)
            / total_batches,
            2,
        )
        if total_batches > 0
        else 0
    )

    average_completion = (
        round(
            sum(b.completion_percentage or 0 for b in batches)
            / total_batches,
            2,
        )
        if total_batches > 0
        else 0
    )

    average_health = (
        round(
            sum(b.health_score or 0 for b in batches)
            / total_batches,
            2,
        )
        if total_batches > 0
        else 0
    )

    db.close()

    return {
        "total_batches": total_batches,
        "completed_batches": completed_batches,
        "ongoing_batches": ongoing_batches,
        "delayed_batches": delayed_batches,
        "average_attendance": average_attendance,
        "average_completion": average_completion,
        "average_health": average_health,
    }
@app.get("/batch-performance")
def batch_performance(course_name: Optional[str] = None, mentor_name: Optional[str] = None):

    db = SessionLocal()

    try:

        query = db.query(Batch)
        if course_name:
            query = query.filter(Batch.course_name == course_name)
        if mentor_name:
            query = query.filter(Batch.mentor_name == mentor_name)

        batches = query.all()

        data = []

        for batch in batches:
            data.append({
                "id": batch.id,
                "batch_name": batch.batch_name,
                "mentor_name": batch.mentor_name,
                "strength": batch.strength,
                "attendance_percentage": batch.attendance_percentage,
                "completion_percentage": batch.completion_percentage,
                "health_score": batch.health_score,
                "status": batch.status,
            })

        return data

    finally:
        db.close()

@app.get("/learner-analytics")
def learner_analytics(course_name: Optional[str] = None, mentor_name: Optional[str] = None):

    db = SessionLocal()

    try:

        query = db.query(Batch)
        if course_name:
            query = query.filter(Batch.course_name == course_name)
        if mentor_name:
            query = query.filter(Batch.mentor_name == mentor_name)

        batches = query.all()

        total_learners = sum(b.strength or 0 for b in batches)

        active_learners = sum(b.active_learners or 0 for b in batches)

        inactive_learners = sum(b.inactive_learners or 0 for b in batches)

        dropout_count = sum(b.dropout_count or 0 for b in batches)

        average_completion = (
            round(
                sum(b.course_completion or 0 for b in batches)
                / len(batches),
                2,
            )
            if batches
            else 0
        )

        return {
            "total_learners": total_learners,
            "active_learners": active_learners,
            "inactive_learners": inactive_learners,
            "dropout_count": dropout_count,
            "average_completion": average_completion,
        }

    finally:
        db.close()

@app.get("/operations-analytics")
def operations_analytics(course_name: Optional[str] = None, mentor_name: Optional[str] = None):

    db = SessionLocal()

    try:

        query = db.query(OperationsAnalytics)
        if course_name:
            batch_names = [
                b.batch_name for b in db.query(Batch).filter(Batch.course_name == course_name).all()
                if b.batch_name
            ]
            query = query.filter(OperationsAnalytics.batch_name.in_(batch_names))
        if mentor_name:
            query = query.filter(OperationsAnalytics.mentor_name == mentor_name)

        operations = query.all()

        total_projects = len(operations)

        total_sessions = sum(o.total_sessions or 0 for o in operations)

        completed_sessions = sum(o.completed_sessions or 0 for o in operations)

        cancelled_sessions = sum(o.cancelled_sessions or 0 for o in operations)

        average_sla = (
            round(
                sum(o.sla_percentage or 0 for o in operations) / len(operations),
                2,
            )
            if operations else 0
        )

        average_completion = (
            round(
                sum(o.completion_percentage or 0 for o in operations) / len(operations),
                2,
            )
            if operations else 0
        )

        average_mentor_utilization = (
            round(
                sum(o.mentor_utilization or 0 for o in operations) / len(operations),
                2,
            )
            if operations else 0
        )

        average_resource_utilization = (
            round(
                sum(o.resource_utilization or 0 for o in operations) / len(operations),
                2,
            )
            if operations else 0
        )

        average_productivity = (
            round(
                sum(o.productivity_score or 0 for o in operations) / len(operations),
                2,
            )
            if operations else 0
        )

        return {
            "total_projects": total_projects,
            "total_sessions": total_sessions,
            "completed_sessions": completed_sessions,
            "cancelled_sessions": cancelled_sessions,
            "average_sla": average_sla,
            "average_completion": average_completion,
            "average_mentor_utilization": average_mentor_utilization,
            "average_resource_utilization": average_resource_utilization,
            "average_productivity": average_productivity,
        }

    finally:
        db.close()

@app.get("/at-risk-batches")
def at_risk_batches(course_name: Optional[str] = None, mentor_name: Optional[str] = None):

    db = SessionLocal()

    try:

        query = db.query(Batch).filter(Batch.health_score < 70)
        if course_name:
            query = query.filter(Batch.course_name == course_name)
        if mentor_name:
            query = query.filter(Batch.mentor_name == mentor_name)

        batches = query.all()

        data = []

        for batch in batches:
            data.append({
                "batch_name": batch.batch_name,
                "mentor_name": batch.mentor_name,
                "attendance": batch.attendance_percentage,
                "completion": batch.completion_percentage,
                "health": batch.health_score,
                "status": batch.status,
            })

        return data

    finally:
        db.close()

@app.get("/session-trend")
def session_trend():

    data = [
        {
            "month": "Jan",
            "completed": 120,
            "cancelled": 8,
        },
        {
            "month": "Feb",
            "completed": 135,
            "cancelled": 6,
        },
        {
            "month": "Mar",
            "completed": 150,
            "cancelled": 10,
        },
        {
            "month": "Apr",
            "completed": 165,
            "cancelled": 5,
        },
        {
            "month": "May",
            "completed": 180,
            "cancelled": 7,
        },
        {
            "month": "Jun",
            "completed": 195,
            "cancelled": 4,
        },
    ]

    return data

@app.get("/revenue-trend")
def revenue_trend():

    data = [
        {
            "month": "Jan",
            "revenue": 125000,
        },
        {
            "month": "Feb",
            "revenue": 148000,
        },
        {
            "month": "Mar",
            "revenue": 172000,
        },
        {
            "month": "Apr",
            "revenue": 195000,
        },
        {
            "month": "May",
            "revenue": 218000,
        },
        {
            "month": "Jun",
            "revenue": 245000,
        },
    ]

    return data

@app.get("/batch-health-chart")
def batch_health_chart():

    data = [
        {
            "batch": "June",
            "health": 82,
        },
        {
            "batch": "July",
            "health": 76,
        },
        {
            "batch": "August",
            "health": 91,
        },
        {
            "batch": "September",
            "health": 68,
        },
        {
            "batch": "October",
            "health": 88,
        },
    ]

    return data

@app.get("/executive-summary")
def executive_summary(
    course_name: Optional[str] = None,
    mentor_name: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):

    db = SessionLocal()

    try:

        # Projects (not yet tracked per-course; shown as an overall figure)
        total_projects = 12

        # Sessions
        session_query = db.query(SessionModel)
        if course_name:
            session_query = session_query.filter(SessionModel.course_name == course_name)
        if mentor_name:
            session_query = session_query.filter(SessionModel.mentor_name == mentor_name)
        if date_from:
            session_query = session_query.filter(SessionModel.session_date >= date_from)
        if date_to:
            session_query = session_query.filter(SessionModel.session_date <= date_to)
        total_sessions = session_query.count()

        # Batches
        batch_query = db.query(Batch)
        if course_name:
            batch_query = batch_query.filter(Batch.course_name == course_name)
        if mentor_name:
            batch_query = batch_query.filter(Batch.mentor_name == mentor_name)
        batches = batch_query.all()
        total_batches = len(batches)

        # Mentors
        mentor_query = db.query(Mentor)
        if mentor_name:
            mentor_query = mentor_query.filter(Mentor.name == mentor_name)
        elif course_name:
            course_mentor_names = [b.mentor_name for b in batches if b.mentor_name]
            mentor_query = mentor_query.filter(Mentor.name.in_(course_mentor_names))
        total_mentors = mentor_query.count()

        # Learners
        total_learners = sum(
            batch.strength or 0
            for batch in batches
        )

        # Temporary Revenue
        total_revenue = 0

        # Temporary KPIs
        active_issues = 3
        health_score = 92

        return {
            "total_projects": total_projects,
            "total_sessions": total_sessions,
            "total_batches": total_batches,
            "total_mentors": total_mentors,
            "total_learners": total_learners,
            "total_revenue": total_revenue,
            "active_issues": active_issues,
            "health_score": health_score,
        }

    except Exception as e:
        return {
            "error": str(e)
        }

    finally:
        db.close()


@app.get("/placement-summary")
def placement_summary():

    return {
        "eligible_students": 320,
        "placed_students": 185,
        "placement_rate": 57.8,
        "interview_scheduled": 42,
        "offers_received": 210,
        "companies_hiring": 28,
        "average_ctc": 7.2,
        "highest_ctc": 18.5,
    }
@app.get("/placement-status")
def placement_status():

    return [
        {
            "status": "Placed",
            "students": 185,
        },
        {
            "status": "Interview Scheduled",
            "students": 42,
        },
        {
            "status": "Offer Received",
            "students": 25,
        },
        {
            "status": "Preparing",
            "students": 68,
        }
    ]


def _gather_analytics_data(db, batch_name=None, mentor_name=None, session_type=None, date_from=None, date_to=None):
    """Builds every section of the Analytics Dashboard as plain dicts/lists,
    applying the same batch/mentor/session-type/date filters as the frontend.
    Mirrors the client-side aggregation in pages/analytics.js exactly, so the
    exported report always matches what's on screen. Shared by the Excel and
    PDF export endpoints."""

    from datetime import datetime as _dt

    def session_dt(s):
        if not s.session_date:
            return None
        try:
            return _dt.strptime(f"{s.session_date} {s.session_time or '00:00'}", "%Y-%m-%d %H:%M")
        except Exception:
            return None

    def overall_rating(n):
        return (n.instructor_rating + n.doubt_rating + n.website_rating) / 3

    def avg(values):
        values = [v for v in values if v is not None]
        return (sum(values) / len(values)) if values else None

    now = _dt.now()

    # ---- Sessions ----
    session_query = db.query(SessionModel)
    if batch_name:
        session_query = session_query.filter(SessionModel.batch_name == batch_name)
    if mentor_name:
        session_query = session_query.filter(SessionModel.mentor_name == mentor_name)
    if session_type:
        session_query = session_query.filter(SessionModel.session_type == session_type)
    if date_from:
        session_query = session_query.filter(SessionModel.session_date >= date_from)
    if date_to:
        session_query = session_query.filter(SessionModel.session_date <= date_to)
    sessions = session_query.all()

    total_sessions = len(sessions)
    completed_sessions = len([s for s in sessions if s.status == "Completed"])
    cancelled_sessions = len([s for s in sessions if s.status == "Cancelled"])
    upcoming_sessions = len([s for s in sessions if s.status == "Scheduled" and session_dt(s) and session_dt(s) > now])
    rescheduled_sessions = len([s for s in sessions if s.remarks and "Rescheduled from" in s.remarks])

    total_session_minutes = sum(s.duration or 0 for s in sessions)
    sessions_with_duration = [s for s in sessions if (s.duration or 0) > 0]
    avg_duration_hours = (total_session_minutes / len(sessions_with_duration) / 60) if sessions_with_duration else None

    sessions_with_attendance = [s for s in sessions if (s.registered_students or 0) > 0]
    avg_attendance = avg([s.attendance_percentage or 0 for s in sessions_with_attendance])

    sessions_with_feedback = [s for s in sessions if (s.feedback_score or 0) > 0]
    avg_session_rating = avg([s.feedback_score for s in sessions_with_feedback])

    mentor_sla = ((total_sessions - cancelled_sessions) / total_sessions * 100) if total_sessions else None
    completion_rate = (completed_sessions / total_sessions * 100) if total_sessions else 0

    session_issues = [
        {"label": "Sessions without mentor", "value": len([s for s in sessions if not s.mentor_name])},
        {"label": "Low attendance (< 50%)", "value": len([s for s in sessions if (s.registered_students or 0) > 0 and (s.attendance_percentage or 0) < 50])},
        {"label": "Low rating (< 3.0)", "value": len([s for s in sessions if (s.feedback_score or 0) > 0 and s.feedback_score < 3])},
        {"label": "Missing recording", "value": len([s for s in sessions if s.status == "Completed" and not s.recording_link])},
        {"label": "Missing feedback", "value": len([s for s in sessions if s.status == "Completed" and not ((s.feedback_score or 0) > 0)])},
    ]
    session_issues = [i for i in session_issues if i["value"] > 0]

    # ---- Mentors & Batches (master lists, not session-filtered — matches the UI) ----
    mentors = db.query(Mentor).all()
    total_mentors = len(mentors)
    active_mentors = len([m for m in mentors if m.status != "Inactive"])

    batches = db.query(Batch).all()
    total_batches = len(batches)
    active_batches = len([b for b in batches if b.status != "Inactive"])

    # ---- NPS feedback (filtered by batch/mentor/date, matching the UI) ----
    nps_query = db.query(NPSFeedback)
    if batch_name:
        nps_query = nps_query.filter(NPSFeedback.batch_name == batch_name)
    if mentor_name:
        nps_query = nps_query.filter(NPSFeedback.mentor_name == mentor_name)
    nps_all = nps_query.all()
    if date_from:
        nps_all = [n for n in nps_all if not n.created_at or n.created_at.strftime("%Y-%m-%d") >= date_from]
    if date_to:
        nps_all = [n for n in nps_all if not n.created_at or n.created_at.strftime("%Y-%m-%d") <= date_to]

    # ---- Mentor stats ----
    mentor_stats = []
    for m in mentors:
        if mentor_name and m.name != mentor_name:
            continue
        m_sessions = [s for s in sessions if s.mentor_name == m.name]
        hours = sum(s.duration or 0 for s in m_sessions) / 60
        att_list = [s.attendance_percentage or 0 for s in m_sessions if (s.registered_students or 0) > 0]
        m_nps = [n for n in nps_all if n.mentor_name == m.name]
        non_cancelled = len([s for s in m_sessions if s.status != "Cancelled"])
        mentor_stats.append({
            "name": m.name,
            "sessions": len(m_sessions),
            "hours": round(hours, 1),
            "attendance": avg(att_list),
            "teaching": avg([n.instructor_rating for n in m_nps]),
            "doubt": avg([n.doubt_rating for n in m_nps]),
            "overall": avg([overall_rating(n) for n in m_nps]),
            "sla": (non_cancelled / len(m_sessions) * 100) if m_sessions else None,
        })

    sessions_conducted = sum(m["sessions"] for m in mentor_stats)
    total_mentor_hours = round(sum(m["hours"] for m in mentor_stats), 1)
    mentor_avg_rating = avg([m["overall"] for m in mentor_stats])
    mentor_avg_attendance = avg([m["attendance"] for m in mentor_stats])
    mentor_avg_sla = avg([m["sla"] for m in mentor_stats])
    top_mentors_by_sessions = sorted(mentor_stats, key=lambda m: m["sessions"], reverse=True)[:5]

    # ---- Batch stats ----
    batch_stats = []
    for b in batches:
        if batch_name and b.batch_name != batch_name:
            continue
        b_sessions = [s for s in sessions if b.batch_name and s.batch_name == b.batch_name]
        with_reg = [s for s in b_sessions if (s.registered_students or 0) > 0]
        attendance = avg([s.attendance_percentage or 0 for s in with_reg])
        completion_sessions = [s for s in with_reg if s.assignment_given]
        completion = avg([(s.assignment_completed or 0) / s.registered_students * 100 for s in completion_sessions]) if completion_sessions else None
        health = "Not Enough Data" if not with_reg else ("Healthy" if attendance >= 75 else "At Risk")
        b_nps = [n for n in nps_all if n.batch_name == b.batch_name]
        rating = avg([overall_rating(n) for n in b_nps])
        completed_count = len([s for s in b_sessions if s.status == "Completed"])
        all_done = len(b_sessions) > 0 and all(s.status != "Scheduled" for s in b_sessions)
        batch_stats.append({
            "batch_name": b.batch_name or "Untitled",
            "mentor_name": b.mentor_name or "Not Assigned",
            "sessions": len(b_sessions),
            "completed_count": completed_count,
            "attendance": attendance,
            "rating": rating,
            "completion": completion,
            "health": health,
            "status": "Completed" if all_done else ("Inactive" if b.status == "Inactive" else "Active"),
        })

    completed_batches = len([b for b in batch_stats if b["status"] == "Completed"])
    batch_total_sessions = sum(b["sessions"] for b in batch_stats)
    batch_avg_attendance = avg([b["attendance"] for b in batch_stats])
    batch_avg_rating = avg([b["rating"] for b in batch_stats])
    batch_avg_completion = avg([b["completion"] for b in batch_stats])
    health_counts = {"Healthy": 0, "At Risk": 0, "Not Enough Data": 0}
    for b in batch_stats:
        health_counts[b["health"]] += 1

    # ---- Session Feedback Analytics (real, from NPS Form data) ----
    feedback_responses = len(nps_all)
    avg_teaching = avg([n.instructor_rating for n in nps_all])
    avg_doubt = avg([n.doubt_rating for n in nps_all])
    avg_overall_exp = avg([n.website_rating for n in nps_all])
    avg_overall_rating = avg([overall_rating(n) for n in nps_all])
    positive_pct = (len([n for n in nps_all if n.nps_score >= 9]) / feedback_responses * 100) if feedback_responses else 0
    negative_pct = (len([n for n in nps_all if n.nps_score <= 6]) / feedback_responses * 100) if feedback_responses else 0

    negative_candidates = [n for n in nps_all if overall_rating(n) <= 2.5 and n.feedback]
    negative_candidates.sort(key=lambda n: n.created_at or _dt.min, reverse=True)
    recent_negative_feedback = [
        {
            "date": n.created_at.strftime("%Y-%m-%d") if n.created_at else "—",
            "batch_name": n.batch_name,
            "mentor_name": n.mentor_name,
            "rating": round(overall_rating(n), 1),
            "feedback": n.feedback,
        }
        for n in negative_candidates[:5]
    ]

    # ---- Attendance Analytics ----
    total_registrations = sum(s.registered_students or 0 for s in sessions)
    total_attendees = sum(s.attended_students or 0 for s in sessions)
    no_show_rate = ((total_registrations - total_attendees) / total_registrations * 100) if total_registrations else None

    low_attendance_sorted = sorted(
        [s for s in sessions_with_attendance if (s.attendance_percentage or 0) < 60],
        key=lambda s: s.attendance_percentage or 0,
    )
    low_attendance_sessions = [
        {"topic": s.topic or "Untitled", "attendance": s.attendance_percentage}
        for s in low_attendance_sorted[:6]
    ]

    return {
        "executive_summary": {
            "total_sessions": total_sessions,
            "completed_sessions": completed_sessions,
            "cancelled_sessions": cancelled_sessions,
            "upcoming_sessions": upcoming_sessions,
            "total_mentors": total_mentors,
            "active_mentors": active_mentors,
            "total_batches": total_batches,
            "active_batches": active_batches,
            "total_session_hours": round(total_session_minutes / 60, 1),
            "avg_attendance": avg_attendance,
            "avg_session_rating": avg_session_rating,
            "mentor_sla": mentor_sla,
            "completion_rate": completion_rate,
        },
        "session_summary": {
            "total_sessions": total_sessions,
            "completed_sessions": completed_sessions,
            "cancelled_sessions": cancelled_sessions,
            "rescheduled_sessions": rescheduled_sessions,
            "upcoming_sessions": upcoming_sessions,
            "total_hours": round(total_session_minutes / 60, 1),
            "avg_duration_hours": avg_duration_hours,
            "avg_attendance": avg_attendance,
        },
        "session_issues": session_issues,
        "mentor_summary": {
            "total_mentors": total_mentors,
            "active_mentors": active_mentors,
            "sessions_conducted": sessions_conducted,
            "total_mentor_hours": total_mentor_hours,
            "avg_rating": mentor_avg_rating,
            "avg_attendance": mentor_avg_attendance,
            "avg_sla": mentor_avg_sla,
        },
        "mentor_stats": mentor_stats,
        "top_mentors_by_sessions": top_mentors_by_sessions,
        "batch_summary": {
            "total_batches": total_batches,
            "active_batches": active_batches,
            "completed_batches": completed_batches,
            "total_sessions": batch_total_sessions,
            "avg_attendance": batch_avg_attendance,
            "avg_rating": batch_avg_rating,
            "avg_completion": batch_avg_completion,
        },
        "batch_stats": batch_stats,
        "health_counts": health_counts,
        "feedback_summary": {
            "responses": feedback_responses,
            "avg_overall_rating": avg_overall_rating,
            "avg_teaching": avg_teaching,
            "avg_doubt": avg_doubt,
            "avg_overall_experience": avg_overall_exp,
            "positive_pct": positive_pct,
            "negative_pct": negative_pct,
        },
        "recent_negative_feedback": recent_negative_feedback,
        "attendance_summary": {
            "total_registrations": total_registrations,
            "total_attendees": total_attendees,
            "avg_attendance": avg_attendance,
            "avg_duration_hours": avg_duration_hours,
            "no_show_rate": no_show_rate,
        },
        "low_attendance_sessions": low_attendance_sessions,
    }


def _filter_description(batch_name, mentor_name, session_type, date_from, date_to):
    parts = []
    if batch_name:
        parts.append(f"Batch = {batch_name}")
    if mentor_name:
        parts.append(f"Mentor = {mentor_name}")
    if session_type:
        parts.append(f"Type = {session_type}")
    if date_from:
        parts.append(f"From {date_from}")
    if date_to:
        parts.append(f"To {date_to}")
    return "; ".join(parts) if parts else "All data (no filters applied)"


@app.get("/export-analytics")
def export_analytics(
    batch_name: Optional[str] = None,
    mentor_name: Optional[str] = None,
    session_type: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    db = SessionLocal()

    try:
        data = _gather_analytics_data(db, batch_name, mentor_name, session_type, date_from, date_to)

        file_name = "analytics_report.xlsx"

        with pd.ExcelWriter(file_name, engine="openpyxl") as writer:
            pd.DataFrame([data["executive_summary"]]).to_excel(writer, sheet_name="Executive Summary", index=False)
            pd.DataFrame([data["session_summary"]]).to_excel(writer, sheet_name="Session Analytics", index=False)
            pd.DataFrame(data["session_issues"]).to_excel(writer, sheet_name="Session Issues", index=False)
            pd.DataFrame([data["mentor_summary"]]).to_excel(writer, sheet_name="Mentor Analytics", index=False)
            pd.DataFrame(data["mentor_stats"]).to_excel(writer, sheet_name="Mentor Performance", index=False)
            pd.DataFrame([data["batch_summary"]]).to_excel(writer, sheet_name="Batch Analytics", index=False)
            pd.DataFrame(data["batch_stats"]).to_excel(writer, sheet_name="Batch Performance", index=False)
            pd.DataFrame([data["feedback_summary"]]).to_excel(writer, sheet_name="Session Feedback", index=False)
            pd.DataFrame(data["recent_negative_feedback"]).to_excel(writer, sheet_name="Negative Feedback", index=False)
            pd.DataFrame([data["attendance_summary"]]).to_excel(writer, sheet_name="Attendance Analytics", index=False)
            pd.DataFrame(data["low_attendance_sessions"]).to_excel(writer, sheet_name="Low Attendance Sessions", index=False)

        return FileResponse(
            file_name,
            filename=file_name,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    finally:
        db.close()


@app.get("/export-analytics-report")
def export_analytics_report(
    batch_name: Optional[str] = None,
    mentor_name: Optional[str] = None,
    session_type: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    db = SessionLocal()

    try:
        data = _gather_analytics_data(db, batch_name, mentor_name, session_type, date_from, date_to)
        filter_desc = _filter_description(batch_name, mentor_name, session_type, date_from, date_to)

        pdf_path = generate_analytics_report(data, filter_desc)

        return FileResponse(
            path=pdf_path,
            filename="Analytics_Report.pdf",
            media_type="application/pdf",
        )

    finally:
        db.close()


from sqlalchemy import func
from models.zoom_analytics import ZoomAnalytics
from models.webinar_participant import WebinarParticipant
import webinar_operations as webinar_ops
from models.webinar_registration import WebinarRegistration


def compute_webinar_health_status(health_score: float) -> str:
    if health_score >= 85:
        return "Excellent"
    if health_score >= 70:
        return "Good"
    if health_score >= 50:
        return "Needs Improvement"
    return "Poor"


def apply_zoom_analytics_fields(record: ZoomAnalytics, data: ZoomAnalyticsCreate):
    for field, value in data.dict().items():
        setattr(record, field, value)

    record.no_show_learners = max(data.registered_learners - data.attended_learners, 0)
    record.attendance_rate = (
        round(data.attended_learners / data.registered_learners * 100, 1)
        if data.registered_learners else 0
    )
    record.no_show_rate = (
        round(record.no_show_learners / data.registered_learners * 100, 1)
        if data.registered_learners else 0
    )


# ==========================================================
# Webinar Report Data Entry (CRUD)
# ==========================================================

@app.post("/zoom-analytics")
def create_zoom_analytics(data: ZoomAnalyticsCreate, db: Session = Depends(get_db)):

    record = ZoomAnalytics()
    apply_zoom_analytics_fields(record, data)
    record.created_at = datetime.now().isoformat()

    db.add(record)
    db.commit()
    db.refresh(record)

    return {"message": "Webinar Report Saved Successfully", "id": record.id}


@app.get("/zoom-analytics")
def list_zoom_analytics(db: Session = Depends(get_db)):
    return db.query(ZoomAnalytics).all()


@app.put("/zoom-analytics/{record_id}")
def update_zoom_analytics(record_id: int, data: ZoomAnalyticsCreate, db: Session = Depends(get_db)):

    record = db.query(ZoomAnalytics).filter(ZoomAnalytics.id == record_id).first()

    if not record:
        return {"message": "Webinar Report Not Found"}

    apply_zoom_analytics_fields(record, data)

    db.commit()

    return {"message": "Webinar Report Updated Successfully"}


@app.delete("/zoom-analytics/{record_id}")
def delete_zoom_analytics(record_id: int, db: Session = Depends(get_db)):

    record = db.query(ZoomAnalytics).filter(ZoomAnalytics.id == record_id).first()

    if record:
        db.delete(record)
        db.commit()

    return {"message": "Webinar Report Deleted Successfully"}


# ==========================================================
# Webinar List API (Dropdown)
# ==========================================================

@app.get("/webinars")
def webinars(db: Session = Depends(get_db), stats: bool = False):

    # stats=true powers the Webinar List table (registration/attendance/
    # leads/payout columns). Default (false) is the original thin dropdown
    # projection zoom-analytics.js already relies on — unchanged.
    if stats:
        return webinar_ops.list_webinars_with_stats(db)

    data = db.query(ZoomAnalytics).all()

    return [
        {
            "session_id": row.session_id,
            "title": row.webinar_title,
            "mentor": row.mentor_name,
            "date": row.session_date,
            "batch": row.batch_name,
            "course": row.course_name,
        }
        for row in data
    ]


# ==========================================================
# Zoom Summary
# ==========================================================

def _zoom_filtered(query, date=None, mentor=None, title=None):
    """Optional dashboard filters shared by the zoom summary / trend / poll
    endpoints. With no arguments the query is returned untouched, so existing
    callers keep getting all-time data."""
    if date:
        query = query.filter(ZoomAnalytics.session_date == date)
    if mentor:
        query = query.filter(ZoomAnalytics.mentor_name == mentor)
    if title is not None and title != "":
        query = query.filter(ZoomAnalytics.webinar_title == title)
    return query


@app.get("/zoom-summary")
def zoom_summary(date: Optional[str] = None, mentor: Optional[str] = None, title: Optional[str] = None, db: Session = Depends(get_db)):

    row = _zoom_filtered(db.query(
        func.count(ZoomAnalytics.id),
        func.sum(ZoomAnalytics.registered_learners),
        func.sum(ZoomAnalytics.attended_learners),
        func.avg(ZoomAnalytics.attendance_rate),
        func.avg(ZoomAnalytics.average_watch_time),
        func.avg(ZoomAnalytics.engagement_score),
        func.avg(ZoomAnalytics.session_rating),
        func.sum(ZoomAnalytics.recording_views),
        func.avg(ZoomAnalytics.poll_response_rate),
        func.avg(ZoomAnalytics.webinar_health_score),
    ), date, mentor, title).one()

    (
        total_webinars,
        registered,
        attended,
        attendance_rate,
        average_watch_time,
        engagement_score,
        session_rating,
        recording_views,
        poll_response_rate,
        webinar_health_score,
    ) = row

    return {
        "total_webinars": total_webinars or 0,
        "registered_learners": registered or 0,
        "attended_learners": attended or 0,
        "attendance_rate": round(attendance_rate or 0, 1),
        "average_watch_time": round(average_watch_time or 0, 1),
        "engagement_score": round(engagement_score or 0, 1),
        "session_rating": round(session_rating or 0, 1),
        "recording_views": recording_views or 0,
        "poll_response_rate": round(poll_response_rate or 0, 1),
        "webinar_health_score": round(webinar_health_score or 0, 1),
        "webinar_health_status": compute_webinar_health_status(webinar_health_score or 0),
    }


# ==========================================================
# Attendance Trend
# ==========================================================

@app.get("/zoom-attendance-trend")
def zoom_attendance_trend(date: Optional[str] = None, mentor: Optional[str] = None, title: Optional[str] = None, db: Session = Depends(get_db)):

    data = _zoom_filtered(db.query(
        ZoomAnalytics.webinar_title,
        ZoomAnalytics.attendance_rate,
        ZoomAnalytics.registered_learners,
        ZoomAnalytics.attended_learners,
        ZoomAnalytics.session_date,
        ZoomAnalytics.mentor_name,
    ), date, mentor, title).all()

    return [
        {
            "meeting": row.webinar_title,
            "attendance": row.attendance_rate,
            "registered": row.registered_learners or 0,
            "attended": row.attended_learners or 0,
            "date": row.session_date,
            "mentor": row.mentor_name,
        }
        for row in data
    ]


# ==========================================================
# Chat Analytics
# ==========================================================

@app.get("/zoom-chat-analytics")
def zoom_chat_analytics(db: Session = Depends(get_db)):

    data = db.query(
        ZoomAnalytics.webinar_title,
        ZoomAnalytics.total_chat_messages,
        ZoomAnalytics.questions_asked,
        ZoomAnalytics.raised_hands
    ).all()

    return [
        {
            "meeting": row.webinar_title,
            "messages": row.total_chat_messages,
            "questions": row.questions_asked,
            "raised_hands": row.raised_hands,
        }
        for row in data
    ]


# ==========================================================
# Poll Analytics
# ==========================================================

@app.get("/zoom-poll-analytics")
def zoom_poll_analytics(date: Optional[str] = None, mentor: Optional[str] = None, title: Optional[str] = None, db: Session = Depends(get_db)):

    data = _zoom_filtered(db.query(
        ZoomAnalytics.webinar_title,
        ZoomAnalytics.polls_conducted,
        ZoomAnalytics.poll_responses,
        ZoomAnalytics.poll_response_rate,
        ZoomAnalytics.poll_average_rating,
    ), date, mentor, title).all()

    return [
        {
            "meeting": row.webinar_title,
            "polls": row.polls_conducted,
            "responses": row.poll_responses,
            "response_rate": row.poll_response_rate,
            "rating": row.poll_average_rating,
        }
        for row in data
    ]


@app.get("/webinar-analytics/export-excel")
def webinar_analytics_export_excel(date: Optional[str] = None, mentor: Optional[str] = None, title: Optional[str] = None, db: Session = Depends(get_db)):
    """Excel version of the Webinar Analytics dashboard, honouring the same
    Date / Mentor / Webinar filters as the on-screen KPIs and charts."""
    rows = _zoom_filtered(db.query(ZoomAnalytics), date, mentor, title).order_by(ZoomAnalytics.session_date).all()

    def avg(values):
        vals = [v for v in values if v is not None]
        return round(sum(vals) / len(vals), 1) if vals else 0

    health = avg([r.webinar_health_score for r in rows])
    summary_rows = [
        {"Metric": "Total Webinars", "Value": len(rows)},
        {"Metric": "Registered Learners", "Value": sum(r.registered_learners or 0 for r in rows)},
        {"Metric": "Attended Learners", "Value": sum(r.attended_learners or 0 for r in rows)},
        {"Metric": "Attendance Rate %", "Value": avg([r.attendance_rate for r in rows])},
        {"Metric": "Poll Response Rate %", "Value": avg([r.poll_response_rate for r in rows])},
        {"Metric": "Session Rating", "Value": avg([r.session_rating for r in rows])},
        {"Metric": "Health Score", "Value": health},
        {"Metric": "Health Status", "Value": compute_webinar_health_status(health)},
    ]
    filter_rows = [
        {"Filter": "Date", "Value": date or "All Dates"},
        {"Filter": "Mentor", "Value": mentor or "All Mentors"},
        {"Filter": "Webinar", "Value": title if title else "All Webinars"},
        {"Filter": "Generated", "Value": datetime.now().strftime("%d %b %Y %H:%M")},
    ]
    attendance_rows = [{
        "Webinar": r.webinar_title, "Mentor": r.mentor_name, "Date": r.session_date,
        "Registered": r.registered_learners or 0, "Attended": r.attended_learners or 0,
        "Attendance %": r.attendance_rate, "No Shows": r.no_show_learners,
    } for r in rows]
    poll_rows = [{
        "Webinar": r.webinar_title, "Polls Conducted": r.polls_conducted, "Poll Responses": r.poll_responses,
        "Response Rate %": r.poll_response_rate, "Average Rating": r.poll_average_rating,
        "Highest Rated Poll": r.highest_rated_poll,
    } for r in rows]
    health_rows = [{
        "Webinar": r.webinar_title, "Engagement Score": r.engagement_score, "Session Rating": r.session_rating,
        "Health Score": r.webinar_health_score, "Health Status": compute_webinar_health_status(r.webinar_health_score or 0),
    } for r in rows]

    file_name = f"Webinar_Analytics_{datetime.utcnow().strftime('%Y-%m-%d')}.xlsx"

    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(file_name, engine="openpyxl") as writer:
        for sheet_name, sheet_rows in [
            ("Filters Applied", filter_rows), ("Summary", summary_rows), ("Attendance", attendance_rows),
            ("Polls", poll_rows), ("Health", health_rows),
        ]:
            df = pd.DataFrame(sheet_rows) if sheet_rows else pd.DataFrame({"Info": ["No data for the selected filters."]})
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="0B1220")
                cell.font = Font(bold=True, color="F5A623")
            ws.freeze_panes = "A2"
            for idx, col in enumerate(df.columns, start=1):
                longest = max([len(str(col))] + [len(str(v)) for v in df[col].tolist()])
                ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 3, 12), 60)

    return FileResponse(file_name, filename=file_name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/webinar-report/{session_id}")
def webinar_report(session_id: int, db: Session = Depends(get_db)):

    webinar = (
        db.query(ZoomAnalytics)
        .filter(ZoomAnalytics.session_id == session_id)
        .first()
    )

    if webinar is None:
        return {"message": "Webinar not found"}

    return {
    # Webinar Information
    "session_id": webinar.session_id,
    "title": webinar.webinar_title,
    "project": webinar.project_name,
    "course": webinar.course_name,
    "batch": webinar.batch_name,
    "mentor": webinar.mentor_name,
    "mentor_email": webinar.mentor_email,
    "date": webinar.session_date,
    "time": webinar.session_time,
    "duration": webinar.duration,
    "platform": webinar.platform,
    "status": webinar.webinar_status,

    # Attendance
    "registered_learners": webinar.registered_learners,
    "attended_learners": webinar.attended_learners,
    "attendance_rate": webinar.attendance_rate,
    "no_show_learners": webinar.no_show_learners,
    "no_show_rate": webinar.no_show_rate,
    "peak_concurrent_users": webinar.peak_concurrent_users,

    # Watch Time
    "average_watch_time": webinar.average_watch_time,
    "late_joiners": webinar.late_joiners,
    "early_exit_learners": webinar.early_exit_learners,

    # Chat
    "total_chat_messages": webinar.total_chat_messages,
    "learner_messages": webinar.learner_messages,
    "mentor_messages": webinar.mentor_messages,
    "questions_asked": webinar.questions_asked,
    "raised_hands": webinar.raised_hands,
    "emoji_reactions": webinar.emoji_reactions,

    # Q&A
    "questions_answered": webinar.questions_answered,
    "average_response_time": webinar.average_response_time,
    "resolved_questions": webinar.resolved_questions,
    "open_questions": webinar.open_questions,

    # Poll
    "polls_conducted": webinar.polls_conducted,
    "poll_responses": webinar.poll_responses,
    "poll_response_rate": webinar.poll_response_rate,
    "poll_average_rating": webinar.poll_average_rating,
    "highest_rated_poll": webinar.highest_rated_poll,

    # Feedback
    "feedback_submitted": webinar.feedback_submitted,
    "session_rating": webinar.session_rating,
    "mentor_rating": webinar.mentor_rating,
    "content_rating": webinar.content_rating,
    "audio_quality_rating": webinar.audio_quality_rating,
    "video_quality_rating": webinar.video_quality_rating,

    # Speaking
    "mentor_speaking_minutes": webinar.mentor_speaking_minutes,
    "learner_speaking_minutes": webinar.learner_speaking_minutes,
    "qa_duration": webinar.qa_duration,
    "discussion_duration": webinar.discussion_duration,

    # Recording
    "recording_available": webinar.recording_available,
    "recording_views": webinar.recording_views,
    "average_recording_watch_time": webinar.average_recording_watch_time,
    "recording_completion_rate": webinar.recording_completion_rate,

    # AI / Health
    "engagement_score": webinar.engagement_score,
    "webinar_health_score": webinar.webinar_health_score,
    "webinar_health_status": compute_webinar_health_status(webinar.webinar_health_score or 0),
    "learner_satisfaction": webinar.learner_satisfaction,

    # Derived Metrics
    "dropout_rate": round((webinar.early_exit_learners / webinar.attended_learners * 100), 1)
        if webinar.attended_learners else 0,
    "qa_resolution_rate": round((webinar.resolved_questions / webinar.questions_asked * 100), 1)
        if webinar.questions_asked else 0,

    "remarks": webinar.remarks,
}


# ==========================================================
# Webinar Registrations (List of Registered Learners)
# ==========================================================

@app.get("/webinar-registrations/{session_id}")
def webinar_registrations(session_id: int, db: Session = Depends(get_db)):

    registrations = (
        db.query(WebinarRegistration)
        .filter(WebinarRegistration.session_id == session_id)
        .all()
    )

    return [
        {
            "id": r.id,
            "learner_name": r.learner_name,
            "learner_email": r.learner_email,
            "phone": r.phone,
            "registered_at": r.registered_at,
            "attended": r.attended,
            "join_time": r.join_time,
            "leave_time": r.leave_time,
            "attendance_duration_minutes": r.attendance_duration_minutes,
        }
        for r in registrations
    ]


@app.delete("/webinar-registrations/{session_id}")
def delete_webinar_registrations(session_id: int, db: Session = Depends(get_db)):

    deleted = (
        db.query(WebinarRegistration)
        .filter(WebinarRegistration.session_id == session_id)
        .delete()
    )
    db.commit()

    return {"message": f"Deleted {deleted} registration(s) for this webinar.", "deleted": deleted}


@app.delete("/webinar-registrations/{session_id}/attendance")
def delete_webinar_attendance(session_id: int, db: Session = Depends(get_db)):
    """Clears attendance data (join/leave/duration/attended) from every
    learner in this webinar without deleting the learners themselves — used
    to undo a bad Attendance import while keeping the registration list."""

    registrations = (
        db.query(WebinarRegistration)
        .filter(WebinarRegistration.session_id == session_id)
        .all()
    )

    for r in registrations:
        r.attended = False
        r.join_time = None
        r.leave_time = None
        r.attendance_duration_minutes = 0

    db.commit()

    return {
        "message": f"Cleared attendance data for {len(registrations)} learner(s).",
        "cleared": len(registrations),
    }


# Zoom export column headers vary by report type (Registration report vs
# Attendee/Attendance report) and by Zoom account settings, so we normalize
# headers and match against every alias we've seen rather than one fixed name.
ZOOM_COLUMN_ALIASES = {
    "learner_name": ["name", "name_original_name", "attendee_name", "full_name", "user_name"],
    "first_name": ["first_name"],
    "last_name": ["last_name"],
    "learner_email": ["email", "user_email", "email_address"],
    "phone": ["phone", "phone_number"],
    "registered_at": ["registration_time", "approval_time", "registered_at", "date"],
    "join_time": ["join_time", "join_time1"],
    "leave_time": ["leave_time", "leave_time1"],
    "duration": ["duration_minutes", "duration_min", "duration", "time_in_session_minutes"],
    "approval_status": ["approval_status", "status"],
    "attended_flag": ["attended", "attended_session"],
}

# Zoom's Registration Report leads with a one-row "meeting summary" table
# (Topic, ID, Scheduled Time, Duration, # Registrants, # Cancelled/Approved/
# Denied registrants) before the real "Attendee Details" table. It's useful
# on its own — it gives an authoritative registered-learner count straight
# from Zoom — so we pull it out separately from the learner rows.
ZOOM_SUMMARY_ALIASES = {
    "topic": ["topic"],
    "meeting_id": ["id", "webinar_id"],
    "scheduled_time": ["scheduled_time"],
    "duration_minutes": ["duration_minutes", "actual_duration_minutes"],
    "total_registrants": ["registrants"],
    "cancelled_registrants": ["cancelled_registrants"],
    "approved_registrants": ["approved_registrants"],
    "denied_registrants": ["denied_registrants"],
    "total_participants": ["participants", "unique_viewers"],
}

APPROVED_STATUSES = {"approved", "", "nan", "none"}

# Zoom writes these as literal cell text for "no value" depending on report
# type — treat them all as empty so they never get stored or misread as real
# data (e.g. a literal "--" in a Join Time column must not count as attended).
EMPTY_CELL_VALUES = {"", "nan", "none", "--", "-"}


def clean_cell(value) -> str:
    text = str(value).strip() if value is not None else ""
    return "" if text.lower() in EMPTY_CELL_VALUES else text


def normalize_header(col) -> str:
    col = str(col).strip().lower()
    col = re.sub(r"[^a-z0-9]+", "_", col)
    return col.strip("_")


def find_header_row(raw_rows, max_scan=15):
    # Zoom labels the real per-learner table "Attendee Details" across every
    # report type (Registration, Attendance, Attendee). A fuller "Attendee
    # Report" export also has "Host Details" / "Panelist Details" sections
    # ABOVE it whose headers also happen to contain "Email" — scanning for
    # the first email-looking header would grab one of those instead and
    # silently drop the real attendee rows, so the section marker takes
    # priority whenever it's present.
    for i, row in enumerate(raw_rows):
        cells = [normalize_header(c) for c in row]
        if len(cells) <= 2 and any(c == "attendee_details" for c in cells):
            return i + 1

    for i, row in enumerate(raw_rows[:max_scan]):
        cells = [normalize_header(c) for c in row]
        if any(c in ZOOM_COLUMN_ALIASES["learner_email"] for c in cells):
            return i
    return 0


def extract_meeting_summary(raw_rows, max_scan=15):
    for i, row in enumerate(raw_rows[:max_scan]):
        cells = [normalize_header(c) for c in row]
        if "topic" in cells and ("registrants" in cells or "participants" in cells) and i + 1 < len(raw_rows):
            data_row = raw_rows[i + 1]
            summary = {}
            for field, aliases in ZOOM_SUMMARY_ALIASES.items():
                for alias in aliases:
                    if alias in cells:
                        idx = cells.index(alias)
                        if idx < len(data_row):
                            summary[field] = data_row[idx]
                        break
            return summary or None
    return None


def decode_csv_bytes(contents: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return contents.decode(encoding)
        except UnicodeDecodeError:
            continue
    return contents.decode("utf-8", errors="ignore")


def read_import_table(contents: bytes, filename: str):
    """Returns (DataFrame of learner rows, meeting summary dict or None)."""
    is_csv = (filename or "").lower().endswith(".csv")

    if is_csv:
        text = decode_csv_bytes(contents)

        # Zoom's CSV exports open with a few metadata rows (report title,
        # topic, meeting ID, ...) that don't have the same column count as
        # the actual table further down, AND — in the fuller "Attendee
        # Report" export — many data rows themselves have a stray trailing
        # comma (an extra empty field) that others in the same table don't.
        # pandas' CSV readers lock the column count to one row and choke or
        # silently misalign on both, so we skip them entirely: read every
        # line with csv.reader (which just returns whatever cells are on
        # that line, ragged or not) and zip each data row against the header
        # by position ourselves, padding short rows and dropping any cells
        # past the header's width.
        raw_rows = list(csv.reader(io.StringIO(text)))
        header_row = find_header_row(raw_rows)
        summary = extract_meeting_summary(raw_rows)

        headers = [normalize_header(c) for c in raw_rows[header_row]] if header_row < len(raw_rows) else []
        width = len(headers)
        records = []
        for row in raw_rows[header_row + 1:]:
            if not any(cell.strip() for cell in row):
                continue
            padded = (row + [""] * width)[:width]
            records.append(dict(zip(headers, padded)))

        df = pd.DataFrame(records, columns=headers)
        return df, summary

    raw = pd.read_excel(io.BytesIO(contents), header=None, dtype=str)
    header_row = find_header_row(raw.values.tolist())
    summary = extract_meeting_summary(raw.values.tolist())
    df = pd.read_excel(io.BytesIO(contents), header=header_row, dtype=str)
    return df, summary


@app.post("/webinar-registrations/{session_id}/import")
async def import_webinar_registrations(session_id: int, file: UploadFile = File(...), db: Session = Depends(get_db)):

    contents = await file.read()

    df, meeting_summary = read_import_table(contents, file.filename)
    df.columns = [normalize_header(c) for c in df.columns]

    def col_for(field):
        for alias in ZOOM_COLUMN_ALIASES[field]:
            if alias in df.columns:
                return alias
        return None

    email_col = col_for("learner_email")
    if not email_col:
        return {"message": "Could not find an email column in this file.", "imported": 0, "updated": 0}

    name_col = col_for("learner_name")
    first_col = col_for("first_name")
    last_col = col_for("last_name")
    registered_col = col_for("registered_at")
    join_col = col_for("join_time")
    leave_col = col_for("leave_time")
    duration_col = col_for("duration")
    status_col = col_for("approval_status")
    phone_col = col_for("phone")
    attended_col = col_for("attended_flag")

    def parse_dt(value):
        try:
            return date_parser.parse(value)
        except (ValueError, TypeError, OverflowError):
            return None

    # Zoom lists the same person more than once whenever they leave and
    # rejoin (a dropped connection, switching devices, ...) — each occurrence
    # is its own row with its own Join/Leave/Duration. We aggregate those by
    # email before writing anything, summing the minutes attended and taking
    # the earliest join / latest leave, instead of letting the last row seen
    # silently overwrite the ones before it.
    aggregated = {}
    skipped, not_approved = 0, 0

    for _, row in df.iterrows():
        email = clean_cell(row.get(email_col))
        if not email:
            skipped += 1
            continue

        # Only approved registrations are real registered learners for our
        # purposes — a cancelled/denied signup shouldn't inflate the count
        # or show up in the learner list.
        if status_col:
            status = clean_cell(row.get(status_col)).lower()
            if status and status not in APPROVED_STATUSES:
                not_approved += 1
                continue

        if name_col:
            name = clean_cell(row.get(name_col))
        elif first_col:
            name = f"{clean_cell(row.get(first_col))} {clean_cell(row.get(last_col))}".strip()
        else:
            name = ""
        name = " ".join(name.split())

        join_time = clean_cell(row.get(join_col)) if join_col else ""
        leave_time = clean_cell(row.get(leave_col)) if leave_col else ""

        duration_raw = row.get(duration_col) if duration_col else None
        try:
            duration = int(float(duration_raw)) if clean_cell(duration_raw) else 0
        except (ValueError, TypeError):
            duration = 0

        # Some Zoom exports give Join/Leave Time but no separate Duration
        # column — derive it ourselves so attendance length is never blank
        # just because that one column happened to be missing.
        if not duration and join_time and leave_time:
            join_dt, leave_dt = parse_dt(join_time), parse_dt(leave_time)
            if join_dt and leave_dt:
                duration = max(int(round((leave_dt - join_dt).total_seconds() / 60)), 0)

        if attended_col:
            attended = clean_cell(row.get(attended_col)).lower() in ("yes", "true", "1")
        else:
            attended = bool(join_time) or duration > 0

        phone = clean_cell(row.get(phone_col)).lstrip("'") if phone_col else ""
        registered_at = clean_cell(row.get(registered_col)) if registered_col else ""

        entry = aggregated.setdefault(email, {
            "name": "", "phone": "", "registered_at": "",
            "join_time": None, "leave_time": None, "duration": 0, "attended": False,
        })
        entry["name"] = name or entry["name"]
        entry["phone"] = phone or entry["phone"]
        entry["registered_at"] = registered_at or entry["registered_at"]
        entry["attended"] = entry["attended"] or attended
        entry["duration"] += duration

        if join_time:
            existing_dt, new_dt = parse_dt(entry["join_time"]) if entry["join_time"] else None, parse_dt(join_time)
            if not entry["join_time"] or (existing_dt and new_dt and new_dt < existing_dt):
                entry["join_time"] = join_time
        if leave_time:
            existing_dt, new_dt = parse_dt(entry["leave_time"]) if entry["leave_time"] else None, parse_dt(leave_time)
            if not entry["leave_time"] or (existing_dt and new_dt and new_dt > existing_dt):
                entry["leave_time"] = leave_time

    existing = {
        r.learner_email: r
        for r in db.query(WebinarRegistration).filter(WebinarRegistration.session_id == session_id).all()
    }

    imported, updated = 0, 0

    for email, entry in aggregated.items():
        record = existing.get(email)
        if record:
            record.learner_name = entry["name"] or record.learner_name
            if entry["phone"]:
                record.phone = entry["phone"]
            if entry["registered_at"]:
                record.registered_at = entry["registered_at"]
            if entry["join_time"]:
                record.join_time = entry["join_time"]
            if entry["leave_time"]:
                record.leave_time = entry["leave_time"]
            if entry["duration"]:
                record.attendance_duration_minutes = entry["duration"]
            record.attended = record.attended or entry["attended"]
            updated += 1
        else:
            db.add(WebinarRegistration(
                session_id=session_id,
                learner_name=entry["name"] or email.split("@")[0],
                learner_email=email,
                phone=entry["phone"] or None,
                registered_at=entry["registered_at"],
                attended=entry["attended"],
                join_time=entry["join_time"],
                leave_time=entry["leave_time"],
                attendance_duration_minutes=entry["duration"],
            ))
            imported += 1

    db.commit()

    message = f"Imported {imported} new, updated {updated} existing, skipped {skipped} rows."
    if not_approved:
        message += f" Excluded {not_approved} cancelled/denied registration(s)."

    return {
        "message": message,
        "imported": imported,
        "updated": updated,
        "skipped": skipped,
        "not_approved": not_approved,
        "meeting_summary": meeting_summary,
    }


POLL_IDENTIFIER_COLUMNS = {"", "user_name", "email_address", "submitted_date_and_time"}


def _section_marker_row(raw_rows, label, start=0):
    """Find a row whose only non-empty cell matches `label` exactly (case-insensitive)."""
    target = label.strip().lower()
    for i in range(start, len(raw_rows)):
        cells = [c.strip() for c in raw_rows[i]]
        if cells and cells[0].lower() == target and not any(cells[1:]):
            return i
    return None


def compute_poll_health_status(avg_rating: float) -> str:
    if not avg_rating:
        return "No Data"
    return "Good" if avg_rating >= 4.3 else "Poor"


def parse_poll_report(raw_rows):
    """Parses Zoom's Poll Report export: an Overview block, a "Launched Polls"
    summary table, then one response table per poll (named after the poll
    itself) with a dynamic set of rating-question columns (1-5) plus an
    optional free-text Remarks column."""

    overview = {}
    overview_row = _section_marker_row(raw_rows, "Overview")
    if overview_row is not None and overview_row + 2 < len(raw_rows):
        header = [normalize_header(c) for c in raw_rows[overview_row + 1]]
        data = raw_rows[overview_row + 2]
        for field in ("generate_time", "meeting_topic", "meeting_webinar_id", "actual_start_time"):
            if field in header:
                idx = header.index(field)
                if idx < len(data):
                    overview[field] = data[idx]

    launched = []
    launched_row = _section_marker_row(raw_rows, "Launched Polls")
    if launched_row is not None and launched_row + 1 < len(raw_rows):
        header = [normalize_header(c) for c in raw_rows[launched_row + 1]]
        j = launched_row + 2
        while j < len(raw_rows) and any(c.strip() for c in raw_rows[j]):
            data = raw_rows[j]
            launched.append({col: (data[k] if k < len(data) else "") for k, col in enumerate(header)})
            j += 1

    polls = []
    search_from = 0
    for entry in launched:
        poll_name = entry.get("poll_name", "").strip()
        if not poll_name:
            continue

        marker = _section_marker_row(raw_rows, poll_name, start=search_from)
        if marker is None or marker + 1 >= len(raw_rows):
            continue
        search_from = marker + 1

        headers = [normalize_header(c) for c in raw_rows[marker + 1]]
        width = len(headers)

        data_rows = []
        j = marker + 2
        while j < len(raw_rows) and any(c.strip() for c in raw_rows[j][:width]):
            data_rows.append((raw_rows[j] + [""] * width)[:width])
            j += 1

        rating_cols = []
        for idx, col in enumerate(headers):
            if col in POLL_IDENTIFIER_COLUMNS or col == "remarks":
                continue
            values = [r[idx].strip() for r in data_rows]
            non_empty = [v for v in values if v]
            numeric = [v for v in non_empty if v.isdigit() and 1 <= int(v) <= 5]
            if non_empty and len(numeric) >= max(1, len(non_empty) * 0.5):
                rating_cols.append(idx)

        # Business rule: average each question's responses on its own first
        # (e.g. every "teaching style" score across all 176 people), then
        # average those per-question averages together — not one blended
        # average across every individual answer, which would weight
        # respondents differently if some skipped a question.
        question_averages = []
        for idx in rating_cols:
            scores = [int(r[idx].strip()) for r in data_rows if r[idx].strip().isdigit()]
            if scores:
                question_averages.append(round(sum(scores) / len(scores), 2))

        avg_rating = round(sum(question_averages) / len(question_averages), 2) if question_averages else 0

        polls.append({
            "name": poll_name,
            "questions": entry.get("questions", ""),
            "responses": len(data_rows),
            "question_averages": question_averages,
            "average_rating": avg_rating,
        })

    total_responses = sum(p["responses"] for p in polls)
    overall_avg = (
        round(sum(p["average_rating"] * p["responses"] for p in polls) / total_responses, 2)
        if total_responses else 0
    )
    highest = max((p["average_rating"] for p in polls), default=0)

    return {
        "overview": overview,
        "polls": polls,
        "polls_conducted": len(polls),
        "poll_responses": total_responses,
        "poll_average_rating": overall_avg,
        "highest_rated_poll": highest,
    }


@app.post("/webinar-registrations/{session_id}/import-polls")
async def import_webinar_polls(session_id: int, file: UploadFile = File(...)):

    contents = await file.read()
    is_csv = (file.filename or "").lower().endswith(".csv")

    if is_csv:
        text = decode_csv_bytes(contents)
        raw_rows = list(csv.reader(io.StringIO(text)))
    else:
        raw = pd.read_excel(io.BytesIO(contents), header=None, dtype=str)
        raw_rows = [["" if pd.isna(c) else str(c) for c in row] for row in raw.values.tolist()]

    result = parse_poll_report(raw_rows)
    health = compute_poll_health_status(result["poll_average_rating"])

    message = (
        f"Parsed {result['polls_conducted']} poll(s), {result['poll_responses']} total responses. "
        f"Average rating {result['poll_average_rating']}/5 — Poll Health: {health}."
    )

    return {
        "message": message,
        **result,
        "poll_health_status": health,
    }


@app.get("/export-webinar-pdf/{session_id}")
def export_webinar_pdf(session_id: int, db: Session = Depends(get_db)):

    webinar = (
        db.query(ZoomAnalytics)
        .filter(ZoomAnalytics.session_id == session_id)
        .first()
    )

    if webinar is None:
        return {"message": "Webinar not found"}

    pdf_path = generate_webinar_report(webinar)

    return FileResponse(
        path=pdf_path,
        filename=f"webinar_report_{session_id}.pdf",
        media_type="application/pdf",
    )


# ==========================================================
# RESOURCE PORTAL
# ==========================================================

def _log_audit(db, action, resource_id=None, session_id=None, details=None, user=None):
    entry = AuditLog(
        user=user,
        action=action,
        resource_id=resource_id,
        session_id=session_id,
        details=details,
    )
    db.add(entry)
    db.commit()


def _log_download(db, download_type, resource_id=None, session_id=None, user_id=None):
    entry = ResourceDownloadLog(
        resource_id=resource_id,
        session_id=session_id,
        user_id=user_id,
        download_type=download_type,
    )
    db.add(entry)
    db.commit()


def _resolve_mentor_email(db, session=None, mentor_id=None, mentor_name=None):
    if session and session.mentor_email:
        return session.mentor_email

    mentor = None
    if mentor_id:
        mentor = db.query(Mentor).filter(Mentor.id == mentor_id).first()
    if not mentor and mentor_name:
        mentor = db.query(Mentor).filter(Mentor.name == mentor_name).first()

    return mentor.email if mentor else None


def _send_resource_email(
    db,
    recipient,
    email_type,
    subject,
    body,
    resource_requirement_id=None,
    session_id=None,
    mentor_id=None,
    reminder_number=None,
):
    settings = get_settings(db)

    if not settings.email_notifications_enabled:
        db.add(ResourceEmailLog(
            resource_requirement_id=resource_requirement_id,
            session_id=session_id,
            mentor_id=mentor_id,
            email=recipient or "",
            email_type=email_type,
            reminder_number=reminder_number,
            status="Skipped",
            error_message="Email notifications are paused in Settings",
        ))
        db.commit()
        return False

    if not recipient:
        db.add(ResourceEmailLog(
            resource_requirement_id=resource_requirement_id,
            session_id=session_id,
            mentor_id=mentor_id,
            email="",
            email_type=email_type,
            reminder_number=reminder_number,
            status="Failed",
            error_message="No mentor email on file",
        ))
        db.commit()
        _log_audit(db, "Email Failed", session_id=session_id, details=f"{email_type}: no recipient email")
        return False

    success, error = send_generic_email(recipient, subject, body)

    db.add(ResourceEmailLog(
        resource_requirement_id=resource_requirement_id,
        session_id=session_id,
        mentor_id=mentor_id,
        email=recipient,
        email_type=email_type,
        reminder_number=reminder_number,
        status="Sent" if success else "Failed",
        error_message=error,
    ))
    db.commit()

    if not success:
        _log_audit(db, "Email Failed", session_id=session_id, details=f"{email_type}: {error}")

    return success


def _send_reminder_for_requirement(db, requirement, session=None):
    """Shared by the manual 'Send Reminder' endpoint and the scheduler.
    Returns (sent: bool, error: str | None) — error is a business-rule rejection
    (not due for another reminder yet, nothing pending), not an SMTP failure.
    Idempotency is governed by next_reminder_at (interval-based, see below) —
    NOT a once-per-calendar-day cap, so a 2-hour cadence can actually fire more
    than once a day."""
    if requirement.status in ("Complete", "Not Required"):
        return False, "Requirement is not pending"

    now = datetime.utcnow()

    if requirement.next_reminder_at and requirement.next_reminder_at > now:
        return False, "Not due for another reminder yet"

    if session is None:
        session = db.query(SessionModel).filter(SessionModel.id == requirement.session_id).first()

    if not session:
        return False, "Session not found"

    recipient = _resolve_mentor_email(db, session, requirement.mentor_id, requirement.mentor_name)

    settings = get_settings(db)

    reminder_number = (requirement.reminder_count or 0) + 1
    is_final = reminder_number >= settings.max_reminders_before_final

    token = get_or_create_session_token(db, session)
    link = submission_url(session)

    subject, body = reminder_email(session, [requirement], submission_url=link, is_final=is_final)

    success = _send_resource_email(
        db,
        recipient,
        "Final Reminder" if is_final else "Reminder",
        subject,
        body,
        resource_requirement_id=requirement.id,
        session_id=requirement.session_id,
        mentor_id=requirement.mentor_id,
        reminder_number=reminder_number,
    )

    requirement.reminder_count = reminder_number
    requirement.last_reminder_sent_at = now
    requirement.next_reminder_at = now + timedelta(hours=settings.reminder_interval_hours or 2.0)
    db.commit()

    _log_audit(
        db,
        "Reminder Sent",
        session_id=requirement.session_id,
        details=f"{requirement.resource_name} (#{reminder_number})",
    )

    return success, None


def _resource_to_dict(r):
    return {
        "id": r.id,
        "session_id": r.session_id,
        "mentor_id": r.mentor_id,
        "mentor_name": r.mentor_name,
        "batch_name": r.batch_name,
        "course_name": r.course_name,
        "session_topic": r.session_topic,
        "session_date": r.session_date,
        "resource_type": r.resource_type,
        "resource_category": r.resource_category,
        "resource_title": r.resource_title,
        "resource_url": r.resource_url,
        "file_path": r.file_path,
        "file_name": r.file_name,
        "file_size": r.file_size,
        "mime_type": r.mime_type,
        "description": r.description,
        "submitted_at": r.submitted_at,
        "created_at": r.created_at,
        "updated_at": r.updated_at,
        "status": r.status,
        "is_required": r.is_required,
        "due_at": r.due_at,
        "received_at": r.received_at,
        "delay_minutes": r.delay_minutes,
        "delay_hours": r.delay_hours,
        "reminder_count": r.reminder_count,
        "last_reminder_sent_at": r.last_reminder_sent_at,
        "next_reminder_at": r.next_reminder_at,
        "created_by": r.created_by,
        "updated_by": r.updated_by,
    }


def _requirement_to_dict(r):
    return {
        "id": r.id,
        "session_id": r.session_id,
        "mentor_id": r.mentor_id,
        "mentor_name": r.mentor_name,
        "resource_type": r.resource_type,
        "resource_category": r.resource_category,
        "resource_name": r.resource_name,
        "is_required": r.is_required,
        "due_at": r.due_at,
        "status": r.status,
        "received_at": r.received_at,
        "reminder_count": r.reminder_count,
        "last_reminder_sent_at": r.last_reminder_sent_at,
        "next_reminder_at": r.next_reminder_at,
        "created_at": r.created_at,
        "updated_at": r.updated_at,
    }


# ----------------------------------------------------------
# Resources (submitted items)
# ----------------------------------------------------------

def _create_resource_core(
    db,
    *,
    session_id,
    mentor_id,
    mentor_name,
    batch_name,
    course_name,
    session_topic,
    session_date,
    resource_type,
    resource_category,
    resource_title,
    resource_url,
    description,
    is_required,
    created_by,
    file,
):
    """Shared by the ops-facing POST /resources (trusts client-supplied
    session/mentor identity) and the public POST /resources/submit/{token}
    (identity is derived from the token, never from the client)."""
    existing = (
        db.query(Resource)
        .filter(
            Resource.session_id == session_id,
            Resource.mentor_name == mentor_name,
            Resource.resource_type == resource_type,
            Resource.resource_title == resource_title,
        )
        .first()
    )

    if existing:
        return {
            "success": False,
            "message": "This resource has already been submitted.",
            "resource_id": existing.id,
        }

    file_meta = {}
    if file is not None and file.filename:
        file_meta = save_file(file)

    now = datetime.utcnow()

    # Look up a matching pending requirement first, so we know the deadline
    # this submission is being measured against.
    requirement = (
        db.query(ResourceRequirement)
        .filter(
            ResourceRequirement.session_id == session_id,
            ResourceRequirement.resource_type == resource_type,
            ResourceRequirement.status.in_(["Pending", "Overdue"]),
        )
        .first()
    )

    # No requirement configured for this session/type — self-declare one so
    # this submission is still trackable (every completed session should have
    # a trackable resource requirement, whether or not Ops pre-configured it).
    if not requirement and session_id is not None:
        requirement = ResourceRequirement(
            session_id=session_id,
            mentor_id=mentor_id,
            mentor_name=mentor_name,
            resource_type=resource_type,
            resource_category=resource_category,
            resource_name=resource_title,
            is_required=is_required,
            status="Pending",
        )
        db.add(requirement)
        db.commit()
        db.refresh(requirement)
    elif requirement and not requirement.resource_category and resource_category:
        requirement.resource_category = resource_category

    due_at = requirement.due_at if requirement else None
    delay_minutes, delay_hours = compute_delay(due_at, now)

    new_resource = Resource(
        session_id=session_id,
        mentor_id=mentor_id,
        mentor_name=mentor_name,
        batch_name=batch_name,
        course_name=course_name,
        session_topic=session_topic,
        session_date=session_date,
        resource_type=resource_type,
        resource_category=resource_category,
        resource_title=resource_title,
        resource_url=resource_url,
        file_path=file_meta.get("file_path"),
        file_name=file_meta.get("file_name"),
        file_size=file_meta.get("file_size"),
        mime_type=file_meta.get("mime_type"),
        description=description,
        is_required=is_required,
        status="Submitted",
        submitted_at=now,
        due_at=due_at,
        received_at=now,
        delay_minutes=delay_minutes,
        delay_hours=delay_hours,
        created_by=created_by,
    )

    db.add(new_resource)
    db.commit()
    db.refresh(new_resource)

    if requirement:
        requirement.received_at = now
        refresh_requirement_status(requirement, now)
        db.commit()

    _log_audit(
        db,
        "Resource Created",
        resource_id=new_resource.id,
        session_id=session_id,
        user=created_by,
    )

    # Thank-you email — only ever reached after the DB commit above succeeded.
    recipient = _resolve_mentor_email(db, mentor_id=mentor_id, mentor_name=mentor_name)
    confirm_subject, confirm_body = confirmation_email(new_resource)
    _send_resource_email(
        db,
        recipient,
        "Submission Confirmation",
        confirm_subject,
        confirm_body,
        session_id=session_id,
        mentor_id=mentor_id,
    )

    if OPS_NOTIFICATION_EMAIL:
        ops_subject, ops_body = ops_notification_email(new_resource)
        _send_resource_email(
            db,
            OPS_NOTIFICATION_EMAIL,
            "Operations Notification",
            ops_subject,
            ops_body,
            session_id=session_id,
            mentor_id=mentor_id,
        )

    return {
        "success": True,
        "message": "Resource submitted successfully",
        "id": new_resource.id,
    }


@app.post("/resources")
def create_resource(
    session_id: int = Form(None),
    mentor_id: int = Form(None),
    mentor_name: str = Form(None),
    batch_name: str = Form(None),
    course_name: str = Form(None),
    session_topic: str = Form(None),
    session_date: str = Form(None),
    resource_type: str = Form(...),
    resource_category: str = Form(None),
    resource_title: str = Form(...),
    resource_url: str = Form(None),
    description: str = Form(None),
    is_required: bool = Form(True),
    created_by: str = Form(None),
    file: UploadFile = File(None),
):
    db = SessionLocal()

    try:
        return _create_resource_core(
            db,
            session_id=session_id,
            mentor_id=mentor_id,
            mentor_name=mentor_name,
            batch_name=batch_name,
            course_name=course_name,
            session_topic=session_topic,
            session_date=session_date,
            resource_type=resource_type,
            resource_category=resource_category,
            resource_title=resource_title,
            resource_url=resource_url,
            description=description,
            is_required=is_required,
            created_by=created_by,
            file=file,
        )
    finally:
        db.close()


# ----------------------------------------------------------
# Public, token-scoped mentor submission (no login) — the session/mentor
# identity is always derived from the token server-side, never trusted from
# the client, so a mentor can only ever act on the one session the link was
# generated for.
# ----------------------------------------------------------

@app.get("/resources/submit/{token}")
def get_submission_context(token: str, db: Session = Depends(get_db)):
    session = get_session_by_token(db, token)
    if not session:
        return {"success": False, "message": "Invalid or expired link"}

    detail = session_resource_detail(db, session.id)

    return {
        "success": True,
        "session": {
            "id": session.id,
            "topic": session.topic,
            "batch_name": session.batch_name,
            "course_name": session.course_name,
            "mentor_name": session.mentor_name,
            "session_date": session.session_date,
        },
        "pending_requirements": [
            _requirement_to_dict(r) for r in detail["requirements"]
            if r.status in ("Pending", "Overdue")
        ],
        "submitted_resources": [_resource_to_dict(r) for r in detail["resources"]],
    }


@app.post("/resources/submit/{token}")
def submit_resource_via_token(
    token: str,
    resource_type: str = Form(...),
    resource_category: str = Form(None),
    resource_title: str = Form(...),
    resource_url: str = Form(None),
    description: str = Form(None),
    is_required: bool = Form(True),
    file: UploadFile = File(None),
):
    db = SessionLocal()

    try:
        session = get_session_by_token(db, token)
        if not session:
            return {"success": False, "message": "Invalid or expired link"}

        return _create_resource_core(
            db,
            session_id=session.id,
            mentor_id=None,
            mentor_name=session.mentor_name,
            batch_name=session.batch_name,
            course_name=session.course_name,
            session_topic=session.topic,
            session_date=session.session_date,
            resource_type=resource_type,
            resource_category=resource_category,
            resource_title=resource_title,
            resource_url=resource_url,
            description=description,
            is_required=is_required,
            created_by="mentor-link",
            file=file,
        )
    finally:
        db.close()


@app.get("/resources")
def list_resources(
    session_id: int = None,
    mentor_id: int = None,
    mentor_name: str = None,
    batch_name: str = None,
    course_name: str = None,
    resource_type: str = None,
    resource_category: str = None,
    status: str = None,
    search: str = None,
):
    db = SessionLocal()

    try:
        query = db.query(Resource)

        if session_id is not None:
            query = query.filter(Resource.session_id == session_id)
        if mentor_id is not None:
            query = query.filter(Resource.mentor_id == mentor_id)
        if mentor_name:
            query = query.filter(Resource.mentor_name == mentor_name)
        if batch_name:
            query = query.filter(Resource.batch_name == batch_name)
        if course_name:
            query = query.filter(Resource.course_name == course_name)
        if resource_type:
            query = query.filter(Resource.resource_type == resource_type)
        if resource_category:
            query = query.filter(Resource.resource_category == resource_category)
        if status:
            query = query.filter(Resource.status == status)
        if search:
            like = f"%{search}%"
            query = query.filter(
                (Resource.resource_title.ilike(like))
                | (Resource.session_topic.ilike(like))
                | (Resource.mentor_name.ilike(like))
                | (Resource.batch_name.ilike(like))
                | (Resource.course_name.ilike(like))
            )

        resources = query.order_by(Resource.created_at.desc()).all()

        return [_resource_to_dict(r) for r in resources]

    finally:
        db.close()


@app.get("/resources/pending")
def get_pending_resources(mentor_name: str = None, batch_name: str = None):
    db = SessionLocal()

    try:
        rows = list_pending_requirements(db)

        if mentor_name:
            rows = [r for r in rows if r["mentor_name"] == mentor_name]
        if batch_name:
            rows = [r for r in rows if r["batch_name"] == batch_name]

        return rows

    finally:
        db.close()


@app.get("/resources/export")
def export_resources(format: str = "excel"):
    db = SessionLocal()

    try:
        rows = build_export_rows(db)
        df = pd.DataFrame(rows)

        if format == "csv":
            file_name = "resource_report.csv"
            df.to_csv(file_name, index=False)
            media_type = "text/csv"
        else:
            file_name = "resource_report.xlsx"
            df.to_excel(file_name, index=False)
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        return FileResponse(file_name, filename=file_name, media_type=media_type)

    finally:
        db.close()


@app.get("/resources/{resource_id}")
def get_resource(resource_id: int):
    db = SessionLocal()

    try:
        r = db.query(Resource).filter(Resource.id == resource_id).first()

        if not r:
            return {"message": "Resource not found"}

        return _resource_to_dict(r)

    finally:
        db.close()


@app.put("/resources/{resource_id}")
def update_resource(resource_id: int, data: ResourceUpdate):
    db = SessionLocal()

    try:
        r = db.query(Resource).filter(Resource.id == resource_id).first()

        if not r:
            return {"message": "Resource not found"}

        if data.resource_title is not None:
            r.resource_title = data.resource_title
        if data.resource_url is not None:
            r.resource_url = data.resource_url
        if data.description is not None:
            r.description = data.description
        if data.status is not None:
            r.status = data.status

        r.updated_by = data.updated_by
        r.updated_at = datetime.utcnow()

        db.commit()
        db.refresh(r)

        _log_audit(
            db,
            "Resource Updated",
            resource_id=r.id,
            session_id=r.session_id,
            user=data.updated_by,
        )

        return {"message": "Resource updated successfully"}

    finally:
        db.close()


@app.delete("/resources/{resource_id}")
def delete_resource(resource_id: int):
    db = SessionLocal()

    try:
        r = db.query(Resource).filter(Resource.id == resource_id).first()

        if not r:
            return {"message": "Resource not found"}

        session_id = r.session_id
        db.delete(r)
        db.commit()

        _log_audit(db, "Resource Deleted", resource_id=resource_id, session_id=session_id)

        return {"message": "Resource deleted successfully"}

    finally:
        db.close()


# ----------------------------------------------------------
# Resource Requirements (what's required for a session)
# ----------------------------------------------------------

def _create_requirements_for_session(db, session, items, mentor_id=None, mentor_name=None, explicit_due_at=None):
    """items: iterable of dicts with resource_type, resource_category (optional),
    resource_name, is_required (optional, default True). Skips any item whose
    (session, resource_type) already has a requirement. Sends the "Initial
    Resource Request" email (with the session's submission link) if anything
    new was created. Shared by the ops-facing POST /resource-requirements and
    the auto-trigger on session completion."""
    settings = get_settings(db)
    due_at = explicit_due_at or compute_due_at(session, settings)
    mentor_name = mentor_name or session.mentor_name

    created = []
    for item in items:
        resource_type = item["resource_type"]
        existing = (
            db.query(ResourceRequirement)
            .filter(
                ResourceRequirement.session_id == session.id,
                ResourceRequirement.resource_type == resource_type,
            )
            .first()
        )
        if existing:
            continue

        requirement = ResourceRequirement(
            session_id=session.id,
            mentor_id=mentor_id,
            mentor_name=mentor_name,
            resource_type=resource_type,
            resource_category=item.get("resource_category"),
            resource_name=item["resource_name"],
            is_required=item.get("is_required", True),
            due_at=due_at,
            status="Pending",
        )
        db.add(requirement)
        created.append(requirement)

    db.commit()

    for r in created:
        db.refresh(r)
        _log_audit(db, "Requirement Created", session_id=r.session_id, details=r.resource_name)

    if created:
        recipient = _resolve_mentor_email(db, session, mentor_id, mentor_name)
        get_or_create_session_token(db, session)
        link = submission_url(session)
        subject, body = initial_request_email(session, created, submission_url=link)
        _send_resource_email(
            db,
            recipient,
            "Initial Resource Request",
            subject,
            body,
            session_id=session.id,
            mentor_id=mentor_id,
        )

    return created


def _ensure_default_requirements(db, session):
    """Auto-creates the standard resource requirement set the first time a
    session is marked Completed — closes the workflow's own first step,
    which otherwise has no UI or automatic trigger at all."""
    already_configured = (
        db.query(ResourceRequirement).filter(ResourceRequirement.session_id == session.id).first()
    )
    if already_configured:
        return []

    return _create_requirements_for_session(db, session, DEFAULT_REQUIREMENTS)


@app.post("/resource-requirements")
def create_resource_requirements(data: ResourceRequirementCreate):
    db = SessionLocal()

    try:
        session = db.query(SessionModel).filter(SessionModel.id == data.session_id).first()
        if not session:
            return {"message": "Session not found", "created_count": 0, "ids": []}

        items = [
            {
                "resource_type": item.resource_type,
                "resource_category": item.resource_category,
                "resource_name": item.resource_name,
                "is_required": item.is_required,
            }
            for item in data.requirements
        ]

        created = _create_requirements_for_session(
            db,
            session,
            items,
            mentor_id=data.mentor_id,
            mentor_name=data.mentor_name,
            explicit_due_at=data.due_at,
        )

        return {
            "message": f"{len(created)} requirement(s) created",
            "created_count": len(created),
            "ids": [r.id for r in created],
        }

    finally:
        db.close()


@app.get("/resource-requirements")
def list_resource_requirements(
    session_id: int = None,
    mentor_id: int = None,
    mentor_name: str = None,
    resource_category: str = None,
    status: str = None,
):
    db = SessionLocal()

    try:
        query = db.query(ResourceRequirement)

        if session_id is not None:
            query = query.filter(ResourceRequirement.session_id == session_id)
        if mentor_id is not None:
            query = query.filter(ResourceRequirement.mentor_id == mentor_id)
        if mentor_name:
            query = query.filter(ResourceRequirement.mentor_name == mentor_name)
        if resource_category:
            query = query.filter(ResourceRequirement.resource_category == resource_category)
        if status:
            query = query.filter(ResourceRequirement.status == status)

        requirements = query.order_by(ResourceRequirement.due_at.asc()).all()

        return [_requirement_to_dict(r) for r in requirements]

    finally:
        db.close()


@app.get("/resource-requirements/{requirement_id}")
def get_resource_requirement(requirement_id: int):
    db = SessionLocal()

    try:
        r = (
            db.query(ResourceRequirement)
            .filter(ResourceRequirement.id == requirement_id)
            .first()
        )

        if not r:
            return {"message": "Requirement not found"}

        return _requirement_to_dict(r)

    finally:
        db.close()


@app.put("/resource-requirements/{requirement_id}/status")
def update_resource_requirement_status(requirement_id: int, data: ResourceRequirementStatusUpdate):
    db = SessionLocal()

    try:
        r = (
            db.query(ResourceRequirement)
            .filter(ResourceRequirement.id == requirement_id)
            .first()
        )

        if not r:
            return {"message": "Requirement not found"}

        r.status = data.status
        r.updated_at = datetime.utcnow()

        if data.status == "Complete" and not r.received_at:
            r.received_at = datetime.utcnow()

        db.commit()

        if data.status == "Complete":
            action = "Requirement Completed"
        elif data.status == "Not Required":
            action = "Requirement Closed"
        else:
            action = "Requirement Updated"

        _log_audit(
            db,
            action,
            session_id=r.session_id,
            details=f"{r.resource_name} -> {data.status}",
            user=data.updated_by,
        )

        return {"message": "Requirement status updated successfully"}

    finally:
        db.close()


@app.post("/resource-requirements/{requirement_id}/send-reminder")
def send_manual_reminder(requirement_id: int):
    db = SessionLocal()

    try:
        r = (
            db.query(ResourceRequirement)
            .filter(ResourceRequirement.id == requirement_id)
            .first()
        )

        if not r:
            return {"message": "Requirement not found"}

        sent, error = _send_reminder_for_requirement(db, r)

        if error:
            return {"success": False, "message": error}

        return {
            "success": sent,
            "message": "Reminder sent successfully" if sent else "Reminder logged but email delivery failed",
        }

    finally:
        db.close()


@app.delete("/resource-requirements/{requirement_id}")
def delete_resource_requirement(requirement_id: int):
    db = SessionLocal()

    try:
        r = (
            db.query(ResourceRequirement)
            .filter(ResourceRequirement.id == requirement_id)
            .first()
        )

        if not r:
            return {"message": "Requirement not found"}

        session_id = r.session_id
        db.delete(r)
        db.commit()

        _log_audit(db, "Requirement Closed", session_id=session_id, details="deleted")

        return {"message": "Requirement deleted successfully"}

    finally:
        db.close()


# ----------------------------------------------------------
# Resource Tracking (cross-session view)
# ----------------------------------------------------------

@app.get("/resource-tracking")
def resource_tracking(
    mentor_name: str = None,
    batch_name: str = None,
    course_name: str = None,
    status: str = None,
):
    db = SessionLocal()

    try:
        rows = build_tracking_table(db)

        if mentor_name:
            rows = [r for r in rows if r["mentor_name"] == mentor_name]
        if batch_name:
            rows = [r for r in rows if r["batch_name"] == batch_name]
        if course_name:
            rows = [r for r in rows if r["course_name"] == course_name]
        if status:
            rows = [r for r in rows if r["status"] == status]

        return rows

    finally:
        db.close()


@app.get("/resource-tracking/{session_id}")
def resource_tracking_detail(session_id: int):
    db = SessionLocal()

    try:
        detail = session_resource_detail(db, session_id)

        if not detail:
            return {"message": "Session not found"}

        session = detail["session"]

        return {
            "session": {
                "id": session.id,
                "topic": session.topic,
                "course_name": session.course_name,
                "batch_name": session.batch_name,
                "mentor_name": session.mentor_name,
                "mentor_email": session.mentor_email,
                "session_date": session.session_date,
                "session_time": session.session_time,
                "status": session.status,
            },
            "status": detail["status"],
            "required_count": detail["required_count"],
            "received_count": detail["received_count"],
            "missing_count": detail["missing_count"],
            "missing_resources": detail["missing_resources"],
            "due_at": detail["due_at"],
            "received_at": detail["received_at"],
            "delay_hours": detail["delay_hours"],
            "reminder_count": detail["reminder_count"],
            "requirements": [_requirement_to_dict(r) for r in detail["requirements"]],
            "resources": [_resource_to_dict(r) for r in detail["resources"]],
            "emails": [
                {
                    "id": e.id,
                    "email": e.email,
                    "email_type": e.email_type,
                    "sent_at": e.sent_at,
                    "reminder_number": e.reminder_number,
                    "status": e.status,
                    "error_message": e.error_message,
                }
                for e in detail["emails"]
            ],
        }

    finally:
        db.close()


# ----------------------------------------------------------
# Resource Download / Preview / LMS Package
# ----------------------------------------------------------

@app.get("/resources/{resource_id}/download")
def download_resource(resource_id: int, user_id: str = None):
    db = SessionLocal()

    try:
        r = db.query(Resource).filter(Resource.id == resource_id).first()

        if not r:
            return {"message": "Resource not found"}

        if not r.file_path or not os.path.exists(r.file_path):
            return {"message": "This resource has no downloadable file — open the URL instead."}

        _log_download(db, "Individual", resource_id=r.id, session_id=r.session_id, user_id=user_id)
        _log_audit(db, "Resource Downloaded", resource_id=r.id, session_id=r.session_id, user=user_id)

        return FileResponse(
            path=r.file_path,
            filename=r.file_name or os.path.basename(r.file_path),
            media_type=r.mime_type or "application/octet-stream",
        )

    finally:
        db.close()


@app.get("/resources/{resource_id}/preview")
def preview_resource(resource_id: int, user_id: str = None):
    db = SessionLocal()

    try:
        r = db.query(Resource).filter(Resource.id == resource_id).first()

        if not r:
            return {"message": "Resource not found"}

        if not r.file_path or not os.path.exists(r.file_path):
            return {"message": "This resource has no previewable file — open the URL instead."}

        _log_download(db, "Preview", resource_id=r.id, session_id=r.session_id, user_id=user_id)

        return FileResponse(
            path=r.file_path,
            media_type=r.mime_type or "application/octet-stream",
            headers={"Content-Disposition": f'inline; filename="{r.file_name}"'},
        )

    finally:
        db.close()


@app.get("/resources/{resource_id}/download-lms")
def download_resource_for_lms(resource_id: int, user_id: str = None):
    db = SessionLocal()

    try:
        r = db.query(Resource).filter(Resource.id == resource_id).first()

        if not r:
            return {"message": "Resource not found"}

        if not r.file_path or not os.path.exists(r.file_path):
            return {"message": "This resource has no downloadable file — open the URL instead."}

        _log_download(db, "Individual", resource_id=r.id, session_id=r.session_id, user_id=user_id)
        _log_audit(db, "Resource Downloaded", resource_id=r.id, session_id=r.session_id, user=user_id, details="LMS filename")

        return FileResponse(
            path=r.file_path,
            filename=lms_filename(r),
            media_type=r.mime_type or "application/octet-stream",
        )

    finally:
        db.close()


@app.get("/sessions/{session_id}/resources/lms-package")
def download_lms_package(session_id: int, user_id: str = None):
    db = SessionLocal()

    try:
        package = build_lms_package(db, session_id)

        if not package:
            return {"message": "No resources found for this session."}

        zip_path, zip_name = package

        _log_download(db, "LMS Package", session_id=session_id, user_id=user_id)
        _log_audit(db, "LMS Package Downloaded", session_id=session_id, user=user_id)

        return FileResponse(
            path=zip_path,
            filename=zip_name,
            media_type="application/zip",
        )

    finally:
        db.close()


# ----------------------------------------------------------
# Reminder Scheduler
# ----------------------------------------------------------

@app.get("/resource-scheduler/status")
def resource_scheduler_status(db: Session = Depends(get_db)):
    settings = get_settings(db)
    return {
        "enabled": settings.reminder_scheduler_enabled,
        "interval_hours": REMINDER_INTERVAL_HOURS,
        "reminder_interval_hours": settings.reminder_interval_hours,
        "reminder_window_start_hour": settings.reminder_window_start_hour,
        "reminder_window_end_hour": settings.reminder_window_end_hour,
        "reminder_timezone": settings.reminder_timezone,
        "weekend_deadline_enabled": settings.weekend_deadline_enabled,
    }


@app.post("/resource-scheduler/run")
def resource_scheduler_run(dry_run: bool = True):
    # respect_toggle=False: a manual/ops-triggered run always executes,
    # bypassing both the enabled toggle and the reminder time window — those
    # only gate the automatic background tick.
    return run_reminder_check(dry_run=dry_run, respect_toggle=False)


# ----------------------------------------------------------
# Settings
# ----------------------------------------------------------

@app.get("/settings")
def read_settings(db: Session = Depends(get_db)):
    return settings_to_dict(get_settings(db))


@app.put("/settings")
def update_settings(data: AppSettingsUpdate, db: Session = Depends(get_db)):
    settings = get_settings(db)

    for field, value in data.dict(exclude_unset=True, exclude={"updated_by"}).items():
        setattr(settings, field, value)

    db.commit()
    db.refresh(settings)
    _log_audit(db, "Settings Updated", user=data.updated_by)

    return settings_to_dict(settings)


# ----------------------------------------------------------
# Resource Analytics
# ----------------------------------------------------------

@app.get("/resource-analytics")
def resource_analytics():
    db = SessionLocal()

    try:
        return compute_resource_analytics(db)

    finally:
        db.close()


@app.get("/resource-analytics/mentors")
def resource_analytics_mentors():
    db = SessionLocal()

    try:
        return compute_mentor_performance(db)

    finally:
        db.close()


@app.get("/resource-analytics/at-risk-mentors")
def resource_analytics_at_risk_mentors():
    db = SessionLocal()

    try:
        return compute_at_risk_mentors(db)

    finally:
        db.close()


@app.get("/resource-analytics/mentor-heatmap")
def resource_analytics_mentor_heatmap(weeks: int = 4):
    db = SessionLocal()

    try:
        return compute_mentor_heatmap(db, weeks=weeks)

    finally:
        db.close()


# ----------------------------------------------------------
# Mentor 360 — Mentor Business Performance
# ----------------------------------------------------------

def _mentor_360_filter_description(course_name, batch_name, mentor_name, date_from, date_to, classification=None, risk=None):
    parts = []
    if course_name:
        parts.append(f"Course = {course_name}")
    if batch_name:
        parts.append(f"Batch = {batch_name}")
    if mentor_name:
        parts.append(f"Mentor = {mentor_name}")
    if classification:
        parts.append(f"Performance = {classification}")
    if risk:
        parts.append(f"Risk = {risk}")
    if date_from:
        parts.append(f"From {date_from}")
    if date_to:
        parts.append(f"To {date_to}")
    return "; ".join(parts) if parts else "All data (no filters applied)"


# Same dimension list and labels the Mentor 360 page uses for its
# "Average Score by Dimension" chart, so exports match the portal.
_MENTOR_360_DIMENSIONS = [
    ("Delivery", "delivery_performance"),
    ("Attendance", "attendance_engagement"),
    ("Learner Exp.", "learner_experience"),
    ("Quality", "session_quality"),
    ("Resources", "resource_compliance"),
    ("Reliability", "reliability"),
    ("Productivity", "productivity"),
    ("Cost Eff.", "cost_efficiency"),
]


def _mentor_360_quadrant(x, y):
    # Thresholds mirror PerformanceMatrix.jsx on the frontend.
    if x >= 75 and y >= 75:
        return "Top Performer"
    if x >= 75:
        return "Quality Concern"
    if y >= 75:
        return "Operational Concern"
    return "Critical"


def _mentor_360_report_analytics(rows):
    """Everything the Mentor 360 page charts, computed from the same filtered
    rows so the PDF/Excel exports show exactly what is on screen."""
    classification_colors = {"Excellent": "#16a34a", "Strong Performer": "#2563eb", "Needs Attention": "#f59e0b", "At Risk": "#ea580c", "Critical": "#dc2626"}
    risk_colors = {"Low": "#16a34a", "Medium": "#f59e0b", "High": "#ea580c", "Critical": "#dc2626"}

    classification = [
        {"label": label, "count": sum(1 for r in rows if r["classification"] == label), "color": color}
        for label, color in classification_colors.items()
    ]
    risk = [
        {"label": label, "count": sum(1 for r in rows if r["risk"] == label), "color": color}
        for label, color in risk_colors.items()
    ]

    dimension_averages = []
    for label, key in _MENTOR_360_DIMENSIONS:
        scores = [r[key]["score"] for r in rows if r.get(key) and isinstance(r[key].get("score"), (int, float))]
        if scores:
            dimension_averages.append({"name": label, "value": round(sum(scores) / len(scores), 1)})

    matrix = []
    for r in rows:
        d, le = r.get("delivery_performance"), r.get("learner_experience")
        x = d["score"] if d and isinstance(d.get("score"), (int, float)) else None
        y = le["score"] if le and isinstance(le.get("score"), (int, float)) else None
        if x is None or y is None:
            continue
        served = (r.get("productivity") or {}).get("learners_served") or 1
        matrix.append({"mentor_name": r["mentor_name"], "x": x, "y": y, "learners_served": served, "quadrant": _mentor_360_quadrant(x, y)})

    return {"classification": classification, "risk": risk, "dimension_averages": dimension_averages, "matrix": matrix}


def _mentor_360_scorecard_filtered(db, course_name, batch_name, mentor_name, date_from, date_to, classification=None, risk=None):
    rows = get_mentor_scorecard(
        db, mentor_name=mentor_name, course_name=course_name, batch_name=batch_name,
        date_from=date_from, date_to=date_to,
    )
    if classification:
        rows = [r for r in rows if r["classification"] == classification]
    if risk:
        rows = [r for r in rows if r["risk"] == risk]
    return rows


def _mentor_360_executive_summary(rows):
    scored = [r for r in rows if isinstance(r["overall_score"], (int, float))]
    avg_score = round(sum(r["overall_score"] for r in scored) / len(scored), 1) if scored else "N/A"

    by_classification = {"Excellent": 0, "Strong Performer": 0, "Needs Attention": 0, "At Risk": 0, "Critical": 0}
    for r in rows:
        if r["classification"] in by_classification:
            by_classification[r["classification"]] += 1

    return {
        "total_mentors": len(rows),
        "average_score": avg_score,
        "excellent": by_classification["Excellent"],
        "strong_performer": by_classification["Strong Performer"],
        "needs_attention": by_classification["Needs Attention"],
        "at_risk": by_classification["At Risk"],
        "critical": by_classification["Critical"],
    }


@app.get("/mentor-360/config")
def mentor_360_config():
    """Read-only — surfaces the scoring constants (currently plain Python
    config, not DB-backed) for display on the Settings page."""
    return {
        "dimension_weights": DIMENSION_WEIGHTS,
        "classification_bands": [{"min_score": t, "label": l} for t, l in CLASSIFICATION_BANDS],
    }


@app.get("/mentor-360/dashboard")
def mentor_360_dashboard(
    course_name: Optional[str] = None,
    batch_name: Optional[str] = None,
    mentor_name: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    classification: Optional[str] = None,
    risk: Optional[str] = None,
):
    db = SessionLocal()

    try:
        rows = _mentor_360_scorecard_filtered(db, course_name, batch_name, mentor_name, date_from, date_to, classification, risk)
        return {
            "kpis": _mentor_360_executive_summary(rows),
            "mentors": rows,
        }
    finally:
        db.close()


@app.get("/mentor-360/scorecard")
def mentor_360_scorecard(
    course_name: Optional[str] = None,
    batch_name: Optional[str] = None,
    mentor_name: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    classification: Optional[str] = None,
    risk: Optional[str] = None,
):
    db = SessionLocal()

    try:
        return _mentor_360_scorecard_filtered(db, course_name, batch_name, mentor_name, date_from, date_to, classification, risk)
    finally:
        db.close()


@app.get("/mentor-360/at-risk")
def mentor_360_at_risk(
    course_name: Optional[str] = None,
    batch_name: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    db = SessionLocal()

    try:
        rows = get_mentor_scorecard(db, course_name=course_name, batch_name=batch_name, date_from=date_from, date_to=date_to)
        return [r for r in rows if r["risk"] in ("High", "Critical")]
    finally:
        db.close()


@app.get("/mentor-360/export-excel")
def mentor_360_export_excel(
    course_name: Optional[str] = None,
    batch_name: Optional[str] = None,
    mentor_name: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    classification: Optional[str] = None,
    risk: Optional[str] = None,
):
    db = SessionLocal()

    try:
        rows = _mentor_360_scorecard_filtered(db, course_name, batch_name, mentor_name, date_from, date_to, classification, risk)
        summary = _mentor_360_executive_summary(rows)

        def g(row, dim, key, default="N/A"):
            d = row.get(dim)
            return d.get(key, default) if d else default

        scorecard_rows = [{
            "Mentor": r["mentor_name"],
            "Overall Score": r["overall_score"],
            "Delivery Score": g(r, "delivery_performance", "score"),
            "Learner Experience": g(r, "learner_experience", "score"),
            "Quality": g(r, "session_quality", "score"),
            "Reliability": g(r, "reliability", "score"),
            "Resource Compliance": g(r, "resource_compliance", "score"),
            "Attendance": g(r, "attendance_engagement", "score"),
            "Productivity": g(r, "productivity", "score"),
            "Cost Efficiency": g(r, "cost_efficiency", "score"),
            "Risk": r["risk"],
            "Classification": r["classification"],
        } for r in rows]

        delivery_rows = [{
            "Mentor": r["mentor_name"],
            "Scheduled": g(r, "delivery_performance", "scheduled"),
            "Completed": g(r, "delivery_performance", "completed"),
            "Cancelled": g(r, "delivery_performance", "cancelled"),
            "Completion %": g(r, "delivery_performance", "completion_percent"),
            "Cancellation %": g(r, "delivery_performance", "cancellation_percent"),
        } for r in rows if r.get("delivery_performance")]

        learner_rows = [{
            "Mentor": r["mentor_name"],
            "Avg Instructor Rating": g(r, "learner_experience", "avg_instructor_rating"),
            "Avg Doubt Rating": g(r, "learner_experience", "avg_doubt_rating"),
            "NPS": g(r, "learner_experience", "nps_score"),
            "Feedback Count": g(r, "learner_experience", "feedback_count"),
        } for r in rows if r.get("learner_experience")]

        resource_rows = [{
            "Mentor": r["mentor_name"],
            "Required": g(r, "resource_compliance", "required"),
            "Received": g(r, "resource_compliance", "received"),
            "Pending": g(r, "resource_compliance", "pending"),
            "Delayed": g(r, "resource_compliance", "delayed"),
            "On-Time %": g(r, "resource_compliance", "on_time_percent"),
            "Avg Delay (hrs)": g(r, "resource_compliance", "avg_delay_hours"),
        } for r in rows if r.get("resource_compliance")]

        productivity_cost_rows = [{
            "Mentor": r["mentor_name"],
            "Sessions Delivered": g(r, "productivity", "sessions_delivered"),
            "Learners Served": g(r, "productivity", "learners_served"),
            "Batches Served": g(r, "productivity", "batches_served"),
            "Cost / Session": g(r, "cost_efficiency", "cost_per_session"),
            "Cost / Hour": g(r, "cost_efficiency", "cost_per_hour"),
            "Total Cost": g(r, "cost_efficiency", "total_cost"),
        } for r in rows if r.get("productivity") or r.get("cost_efficiency")]

        analytics = _mentor_360_report_analytics(rows)
        filter_desc = _mentor_360_filter_description(course_name, batch_name, mentor_name, date_from, date_to, classification, risk)

        applied = [
            ("Course", course_name), ("Batch", batch_name), ("Mentor", mentor_name),
            ("Performance", classification), ("Risk", risk), ("Date From", date_from), ("Date To", date_to),
        ]
        filters_rows = [{"Filter": k, "Value": v or "All"} for k, v in applied]
        filters_rows.append({"Filter": "Scope", "Value": filter_desc})
        filters_rows.append({"Filter": "Generated", "Value": datetime.now().strftime("%d %b %Y %H:%M")})

        # Same tiles, in the same order, as the KPI row on the Mentor 360 page.
        summary_rows = [
            {"Metric": "Total Mentors", "Value": summary["total_mentors"]},
            {"Metric": "Avg Business Score", "Value": summary["average_score"]},
            {"Metric": "Excellent", "Value": summary["excellent"]},
            {"Metric": "Strong Performers", "Value": summary["strong_performer"]},
            {"Metric": "Needs Attention", "Value": summary["needs_attention"]},
            {"Metric": "At Risk", "Value": summary["at_risk"]},
            {"Metric": "Critical", "Value": summary["critical"]},
        ]

        classification_rows = [{"Performance Classification": c["label"], "Mentors": c["count"]} for c in analytics["classification"]]
        risk_rows = [{"Risk Level": c["label"], "Mentors": c["count"]} for c in analytics["risk"]]
        dimension_avg_rows = [{"Dimension": d["name"], "Average Score": d["value"]} for d in analytics["dimension_averages"]]
        matrix_rows = [{
            "Mentor": m["mentor_name"],
            "Delivery Performance (X)": m["x"],
            "Learner Experience (Y)": m["y"],
            "Learners Served (size)": m["learners_served"],
            "Quadrant": m["quadrant"],
        } for m in analytics["matrix"]]

        attendance_rows = [{
            "Mentor": r["mentor_name"],
            "Avg Attendance %": g(r, "attendance_engagement", "avg_attendance_percent"),
            "Registered Learners": g(r, "attendance_engagement", "registered_learners"),
            "Attended Learners": g(r, "attendance_engagement", "attended_learners"),
            "Low Attendance Sessions": g(r, "attendance_engagement", "low_attendance_sessions"),
            "Avg Session Feedback": g(r, "session_quality", "avg_session_feedback_score"),
            "Quality Score": g(r, "session_quality", "score"),
        } for r in rows if r.get("attendance_engagement") or r.get("session_quality")]

        file_name = f"Mentor_Performance_Report_{datetime.utcnow().strftime('%Y-%m-%d')}.xlsx"

        sheets = [
            ("Filters Applied", filters_rows),
            ("Executive Summary", summary_rows),
            ("Performance Classification", classification_rows),
            ("Risk Distribution", risk_rows),
            ("Dimension Averages", dimension_avg_rows),
            ("Performance Matrix", matrix_rows),
            ("Mentor Scorecard", scorecard_rows),
            ("Delivery Performance", delivery_rows),
            ("Attendance & Quality", attendance_rows),
            ("Learner Experience", learner_rows),
            ("Resource Compliance", resource_rows),
            ("Productivity & Cost", productivity_cost_rows),
        ]

        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        header_fill = PatternFill("solid", fgColor="0B1220")
        header_font = Font(bold=True, color="F5A623")

        with pd.ExcelWriter(file_name, engine="openpyxl") as writer:
            for sheet_name, sheet_rows in sheets:
                df = pd.DataFrame(sheet_rows) if sheet_rows else pd.DataFrame({"Info": ["No data for the selected filters."]})
                df.to_excel(writer, sheet_name=sheet_name, index=False)
                ws = writer.sheets[sheet_name]
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(vertical="center")
                ws.freeze_panes = "A2"
                for idx, col in enumerate(df.columns, start=1):
                    longest = max([len(str(col))] + [len(str(v)) for v in df[col].tolist()])
                    ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 3, 12), 60)

        return FileResponse(
            file_name,
            filename=file_name,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    finally:
        db.close()


@app.get("/mentor-360/export-pdf")
def mentor_360_export_pdf(
    course_name: Optional[str] = None,
    batch_name: Optional[str] = None,
    mentor_name: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    classification: Optional[str] = None,
    risk: Optional[str] = None,
):
    db = SessionLocal()

    try:
        rows = _mentor_360_scorecard_filtered(db, course_name, batch_name, mentor_name, date_from, date_to, classification, risk)
        data = {
            "executive_summary": _mentor_360_executive_summary(rows),
            "scorecard": rows,
            "analytics": _mentor_360_report_analytics(rows),
        }
        filter_desc = _mentor_360_filter_description(course_name, batch_name, mentor_name, date_from, date_to, classification, risk)

        pdf_path = generate_mentor_performance_report(data, filter_desc)

        return FileResponse(
            path=pdf_path,
            filename="Mentor_Business_Performance_Report.pdf",
            media_type="application/pdf",
        )
    finally:
        db.close()


# NOTE: the two /mentor-360/{mentor_name} routes below MUST stay after every
# literal /mentor-360/... route above (export-excel, export-pdf, at-risk, etc.)
# — FastAPI matches path routes in registration order, and {mentor_name} would
# otherwise swallow "export-excel"/"export-pdf" as a mentor name.
@app.get("/mentor-360/{mentor_name}")
def mentor_360_detail(
    mentor_name: str,
    course_name: Optional[str] = None,
    batch_name: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    db = SessionLocal()

    try:
        rows = get_mentor_scorecard(db, mentor_name=mentor_name, course_name=course_name, batch_name=batch_name, date_from=date_from, date_to=date_to)
        if not rows:
            return {"success": False, "message": "No data found for this mentor in the selected scope."}
        return {"success": True, "mentor": rows[0]}
    finally:
        db.close()


@app.get("/mentor-360/{mentor_name}/trends")
def mentor_360_trends(mentor_name: str, months: int = 6):
    db = SessionLocal()

    try:
        return get_mentor_trend(db, mentor_name, months=months)
    finally:
        db.close()


# ----------------------------------------------------------
# Webinar Operations — Phase 1 (DB + backend: CRUD, participants,
# reports, leads, payout). ZoomAnalytics IS the Webinar entity — no
# separate Webinar table. The pre-existing /webinars, /webinar-report/
# {session_id}, /zoom-summary etc. routes above are untouched.
# ----------------------------------------------------------

class WebinarCreate(BaseModel):
    webinar_title: str
    mentor_name: str
    session_date: str
    session_time: Optional[str] = None
    duration: Optional[int] = None
    platform: Optional[str] = None
    webinar_status: Optional[str] = "Scheduled"
    description: Optional[str] = None
    category: Optional[str] = None
    target_audience: Optional[str] = None
    project_name: Optional[str] = None
    batch_name: Optional[str] = None
    course_name: Optional[str] = None


class ParticipantCreate(BaseModel):
    name: str
    email: str
    phone: Optional[str] = None
    source: Optional[str] = None
    course_interest: Optional[str] = None


class ParticipantUpdate(BaseModel):
    attended: Optional[bool] = None
    attendance_duration: Optional[int] = None
    attendance_percentage: Optional[float] = None
    rating: Optional[int] = None
    feedback: Optional[str] = None
    course_interest: Optional[str] = None
    lead_status: Optional[str] = None
    sales_followup_status: Optional[str] = None
    follow_up_date: Optional[datetime] = None


class WebinarReportUpdate(BaseModel):
    registered_learners: Optional[int] = None
    attended_learners: Optional[int] = None
    session_rating: Optional[float] = None
    feedback_submitted: Optional[int] = None
    average_watch_time: Optional[float] = None


def _webinar_validation_error(data: WebinarCreate):
    if data.duration is not None and data.duration <= 0:
        return "Duration must be greater than 0."
    if data.webinar_status and data.webinar_status not in webinar_ops.VALID_STATUSES:
        return f"Status must be one of: {', '.join(webinar_ops.VALID_STATUSES)}"
    return None


@app.get("/webinars/leads")
def webinars_leads(
    mentor_name: Optional[str] = None,
    lead_status: Optional[str] = None,
    course_interest: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    db = SessionLocal()
    try:
        return webinar_ops.list_leads(
            db, mentor_name=mentor_name, lead_status=lead_status,
            course_interest=course_interest, date_from=date_from, date_to=date_to,
        )
    finally:
        db.close()


@app.get("/webinars/lead-profile")
def webinars_lead_profile(email: str):
    db = SessionLocal()
    try:
        profile = webinar_ops.lead_profile(db, email)
        if not profile:
            return {"success": False, "message": "No webinar activity found for this email."}
        return {"success": True, "profile": profile}
    finally:
        db.close()


@app.get("/webinars/mentor-performance/{mentor_name}")
def webinars_mentor_performance(mentor_name: str):
    """Supplementary webinar stats for one mentor — fetched by the existing
    Mentor 360 detail page (/mentor-performance/{mentor}) as an additional
    section, not a competing scoring system."""
    db = SessionLocal()
    try:
        stats = webinar_ops.compute_mentor_webinar_performance(db, mentor_name)
        if not stats:
            return {"success": False, "message": "This mentor has no webinars on record."}
        return {"success": True, "stats": stats}
    finally:
        db.close()


@app.get("/webinars/config")
def webinars_config():
    """Read-only — surfaces the status/lead-status vocabularies (currently
    plain Python constants, not DB-backed) for display on the Settings page."""
    return {
        "webinar_statuses": webinar_ops.VALID_STATUSES,
        "lead_statuses": webinar_ops.VALID_LEAD_STATUSES,
    }


@app.get("/webinars/payouts")
def webinars_payouts_list(mentor_name: Optional[str] = None, payment_status: Optional[str] = None):
    db = SessionLocal()
    try:
        return webinar_ops.list_webinar_payouts(db, mentor_name=mentor_name, payment_status=payment_status)
    finally:
        db.close()


@app.get("/webinars/dashboard")
def webinars_dashboard(
    mentor_name: Optional[str] = None,
    category: Optional[str] = None,
    status: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    db = SessionLocal()
    try:
        return webinar_ops.compute_webinar_dashboard(
            db, mentor_name=mentor_name, category=category, status=status, date_from=date_from, date_to=date_to,
        )
    finally:
        db.close()


@app.get("/webinars/business-insights")
def webinars_business_insights(
    mentor_name: Optional[str] = None,
    category: Optional[str] = None,
    status: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    db = SessionLocal()
    try:
        return webinar_ops.compute_business_insights(
            db, mentor_name=mentor_name, category=category, status=status, date_from=date_from, date_to=date_to,
        )
    finally:
        db.close()


def _webinar_filter_description(mentor_name, category, status, date_from, date_to):
    parts = []
    if mentor_name:
        parts.append(f"Mentor = {mentor_name}")
    if category:
        parts.append(f"Category = {category}")
    if status:
        parts.append(f"Status = {status}")
    if date_from:
        parts.append(f"From {date_from}")
    if date_to:
        parts.append(f"To {date_to}")
    return "; ".join(parts) if parts else "All data (no filters applied)"


@app.get("/webinars/export/excel")
def webinars_export_excel(
    mentor_name: Optional[str] = None,
    category: Optional[str] = None,
    status: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    db = SessionLocal()
    try:
        data = webinar_ops.build_export_data(db, mentor_name, category, status, date_from, date_to)
        file_name = f"Webinar_Operations_Report_{datetime.utcnow().strftime('%Y-%m-%d')}.xlsx"

        with pd.ExcelWriter(file_name, engine="openpyxl") as writer:
            pd.DataFrame([data["summary"]]).to_excel(writer, sheet_name="Executive Summary", index=False)
            pd.DataFrame(data["webinar_rows"]).to_excel(writer, sheet_name="Webinars", index=False)
            pd.DataFrame(data["lead_rows"]).to_excel(writer, sheet_name="Leads", index=False)
            pd.DataFrame(data["payout_rows"]).to_excel(writer, sheet_name="Payouts", index=False)

        return FileResponse(file_name, filename=file_name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    finally:
        db.close()


@app.get("/webinars/export/pdf")
def webinars_export_pdf(
    mentor_name: Optional[str] = None,
    category: Optional[str] = None,
    status: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
):
    db = SessionLocal()
    try:
        data = webinar_ops.build_export_data(db, mentor_name, category, status, date_from, date_to)
        filter_desc = _webinar_filter_description(mentor_name, category, status, date_from, date_to)
        pdf_path = generate_webinar_report_pdf(data, filter_desc)

        return FileResponse(path=pdf_path, filename="Webinar_Operations_Report.pdf", media_type="application/pdf")
    finally:
        db.close()


@app.post("/webinars/import/preview")
async def webinars_import_preview(file: UploadFile = File(...)):
    try:
        file_bytes = await file.read()
        return {"success": True, **webinar_ops.preview_import(file_bytes, file.filename)}
    except Exception as e:
        return {"success": False, "message": f"Could not read this file: {e}"}


@app.post("/webinars")
def webinars_create(data: WebinarCreate):
    db = SessionLocal()
    try:
        error = _webinar_validation_error(data)
        if error:
            return {"success": False, "message": error}

        mentor = db.query(Mentor).filter(Mentor.name == data.mentor_name).first()

        webinar = ZoomAnalytics(
            webinar_title=data.webinar_title,
            description=data.description,
            category=data.category,
            target_audience=data.target_audience,
            project_name=data.project_name,
            batch_name=data.batch_name,
            course_name=data.course_name,
            mentor_name=data.mentor_name,
            mentor_email=mentor.email if mentor else None,
            session_date=data.session_date,
            session_time=data.session_time,
            duration=data.duration,
            platform=data.platform,
            webinar_status=data.webinar_status or "Scheduled",
            created_at=datetime.utcnow().isoformat(),
        )
        db.add(webinar)
        db.commit()
        db.refresh(webinar)

        return {"success": True, "message": "Webinar scheduled", "id": webinar.id}
    finally:
        db.close()


@app.get("/webinars/{webinar_id}")
def webinars_get(webinar_id: int):
    db = SessionLocal()
    try:
        webinar = db.query(ZoomAnalytics).filter(ZoomAnalytics.id == webinar_id).first()
        if not webinar:
            return {"success": False, "message": "Webinar not found"}
        return {"success": True, "webinar": {c.name: getattr(webinar, c.name) for c in ZoomAnalytics.__table__.columns}}
    finally:
        db.close()


@app.patch("/webinars/{webinar_id}")
def webinars_update(webinar_id: int, data: WebinarCreate):
    db = SessionLocal()
    try:
        webinar = db.query(ZoomAnalytics).filter(ZoomAnalytics.id == webinar_id).first()
        if not webinar:
            return {"success": False, "message": "Webinar not found"}

        error = _webinar_validation_error(data)
        if error:
            return {"success": False, "message": error}

        mentor = db.query(Mentor).filter(Mentor.name == data.mentor_name).first()

        webinar.webinar_title = data.webinar_title
        webinar.description = data.description
        webinar.category = data.category
        webinar.target_audience = data.target_audience
        webinar.project_name = data.project_name
        webinar.batch_name = data.batch_name
        webinar.course_name = data.course_name
        webinar.mentor_name = data.mentor_name
        webinar.mentor_email = mentor.email if mentor else webinar.mentor_email
        webinar.session_date = data.session_date
        webinar.session_time = data.session_time
        webinar.duration = data.duration
        webinar.platform = data.platform
        webinar.webinar_status = data.webinar_status or webinar.webinar_status

        db.commit()
        return {"success": True, "message": "Webinar updated"}
    finally:
        db.close()


@app.delete("/webinars/{webinar_id}")
def webinars_delete(webinar_id: int):
    db = SessionLocal()
    try:
        webinar = db.query(ZoomAnalytics).filter(ZoomAnalytics.id == webinar_id).first()
        if not webinar:
            return {"success": False, "message": "Webinar not found"}

        has_participants = db.query(WebinarParticipant).filter(WebinarParticipant.webinar_id == webinar_id).first()
        if has_participants:
            return {"success": False, "message": "This webinar has participant records — mark it Cancelled instead of deleting."}

        db.delete(webinar)
        db.commit()
        return {"success": True, "message": "Webinar deleted"}
    finally:
        db.close()


@app.get("/webinars/{webinar_id}/participants")
def webinars_participants_list(webinar_id: int):
    db = SessionLocal()
    try:
        rows = (
            db.query(WebinarParticipant)
            .filter(WebinarParticipant.webinar_id == webinar_id)
            .order_by(WebinarParticipant.registration_date.desc())
            .all()
        )
        return [{c.name: getattr(r, c.name) for c in WebinarParticipant.__table__.columns} for r in rows]
    finally:
        db.close()


@app.post("/webinars/{webinar_id}/participants")
def webinars_participants_create(webinar_id: int, data: ParticipantCreate):
    db = SessionLocal()
    try:
        webinar = db.query(ZoomAnalytics).filter(ZoomAnalytics.id == webinar_id).first()
        if not webinar:
            return {"success": False, "message": "Webinar not found"}

        return webinar_ops.register_participant(
            db, webinar_id,
            name=data.name, email=data.email, phone=data.phone,
            source=data.source, course_interest=data.course_interest,
        )
    finally:
        db.close()


@app.post("/webinars/{webinar_id}/import/participants")
async def webinars_import_participants(webinar_id: int, file: UploadFile = File(...), column_mapping: str = Form(...)):
    """column_mapping is a JSON string (sent as a form field alongside the
    file) mapping target fields (name/email/phone/...) to the uploaded
    file's actual column names — built from the /webinars/import/preview
    step, so this works with any reasonably-shaped CSV/XLSX, not just one
    hardcoded sheet layout."""
    db = SessionLocal()
    try:
        try:
            mapping = json.loads(column_mapping)
        except json.JSONDecodeError:
            return {"success": False, "message": "Invalid column mapping."}

        file_bytes = await file.read()
        return webinar_ops.commit_participant_import(db, webinar_id, file_bytes, file.filename, mapping)
    except Exception as e:
        return {"success": False, "message": f"Import failed: {e}"}
    finally:
        db.close()


@app.patch("/webinars/{webinar_id}/participants/{participant_id}")
def webinars_participants_update(webinar_id: int, participant_id: int, data: ParticipantUpdate):
    db = SessionLocal()
    try:
        participant = (
            db.query(WebinarParticipant)
            .filter(WebinarParticipant.id == participant_id, WebinarParticipant.webinar_id == webinar_id)
            .first()
        )
        if not participant:
            return {"success": False, "message": "Participant not found"}

        if data.rating is not None and not (1 <= data.rating <= 5):
            return {"success": False, "message": "Rating must be between 1 and 5."}
        if data.lead_status is not None and data.lead_status not in webinar_ops.VALID_LEAD_STATUSES:
            return {"success": False, "message": f"Lead status must be one of: {', '.join(webinar_ops.VALID_LEAD_STATUSES)}"}

        for field, value in data.dict(exclude_unset=True).items():
            setattr(participant, field, value)

        db.commit()
        return {"success": True, "message": "Participant updated"}
    finally:
        db.close()


@app.get("/webinars/{webinar_id}/report")
def webinars_report_get(webinar_id: int):
    db = SessionLocal()
    try:
        report = webinar_ops.compute_webinar_report(db, webinar_id)
        if not report:
            return {"success": False, "message": "Webinar not found"}
        return {"success": True, **report}
    finally:
        db.close()


@app.patch("/webinars/{webinar_id}/report")
def webinars_report_update(webinar_id: int, data: WebinarReportUpdate):
    """Manual report entry path — for webinars with no participant-level
    data, ops can still record the aggregate outcome directly onto the
    ZoomAnalytics row."""
    db = SessionLocal()
    try:
        webinar = db.query(ZoomAnalytics).filter(ZoomAnalytics.id == webinar_id).first()
        if not webinar:
            return {"success": False, "message": "Webinar not found"}

        if data.registered_learners is not None and data.attended_learners is not None:
            if data.attended_learners > data.registered_learners:
                return {"success": False, "message": "Attended cannot exceed registered."}

        for field, value in data.dict(exclude_unset=True).items():
            setattr(webinar, field, value)

        if webinar.registered_learners:
            webinar.attendance_rate = round((webinar.attended_learners or 0) / webinar.registered_learners * 100, 1)

        db.commit()
        return {"success": True, "message": "Report updated"}
    finally:
        db.close()


@app.get("/webinars/{webinar_id}/payout")
def webinars_payout_preview(webinar_id: int):
    db = SessionLocal()
    try:
        webinar = db.query(ZoomAnalytics).filter(ZoomAnalytics.id == webinar_id).first()
        if not webinar:
            return {"success": False, "message": "Webinar not found"}

        mentor = db.query(Mentor).filter(Mentor.name == webinar.mentor_name).first()
        if not mentor or not mentor.hourly_rate:
            return {"success": False, "message": "No hourly rate on file for this mentor."}

        amount = webinar_ops.calculate_payout(mentor.hourly_rate, webinar.duration)
        existing = db.query(Invoice).filter(Invoice.webinar_id == webinar_id, Invoice.source_type == "webinar").first()

        return {
            "success": True,
            "mentor_name": webinar.mentor_name,
            "duration_minutes": webinar.duration,
            "hourly_rate": mentor.hourly_rate,
            "estimated_amount": amount,
            "already_invoiced": existing is not None,
            "invoice_id": existing.id if existing else None,
        }
    finally:
        db.close()


@app.post("/webinars/{webinar_id}/payout")
def webinars_payout_create(webinar_id: int):
    db = SessionLocal()
    try:
        return webinar_ops.create_webinar_payout(db, webinar_id)
    finally:
        db.close()


# ==========================================================
# SESSION REPORTS
# ==========================================================

def _session_report_filters(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    status: Optional[str] = None,
    mentor_name: Optional[str] = None,
    batch_name: Optional[str] = None,
    course_name: Optional[str] = None,
    session_type: Optional[str] = None,
    recording_status: Optional[str] = None,
    report_status: Optional[str] = None,
):
    return {
        "date_from": date_from,
        "date_to": date_to,
        "status": status,
        "mentor_name": mentor_name,
        "batch_name": batch_name,
        "course_name": course_name,
        "session_type": session_type,
        "recording_status": recording_status,
        "report_status": report_status,
    }


@app.get("/session-reports")
def get_session_reports(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    status: Optional[str] = None,
    mentor_name: Optional[str] = None,
    batch_name: Optional[str] = None,
    course_name: Optional[str] = None,
    session_type: Optional[str] = None,
    recording_status: Optional[str] = None,
    report_status: Optional[str] = None,
    search: Optional[str] = None,
    page: int = 1,
    page_size: int = 20,
    sort_by: str = "session_date",
    sort_dir: str = "desc",
    db: Session = Depends(get_db),
):
    filters = _session_report_filters(
        date_from, date_to, status, mentor_name, batch_name,
        course_name, session_type, recording_status, report_status,
    )
    return session_reports.list_session_reports(db, filters, page, page_size, sort_by, sort_dir, search)


@app.get("/session-reports/summary")
def get_session_reports_summary(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    status: Optional[str] = None,
    mentor_name: Optional[str] = None,
    batch_name: Optional[str] = None,
    course_name: Optional[str] = None,
    session_type: Optional[str] = None,
    recording_status: Optional[str] = None,
    report_status: Optional[str] = None,
    db: Session = Depends(get_db),
):
    filters = _session_report_filters(
        date_from, date_to, status, mentor_name, batch_name,
        course_name, session_type, recording_status, report_status,
    )
    return session_reports.session_reports_summary(db, filters)


@app.get("/session-reports/export")
def export_session_reports(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    status: Optional[str] = None,
    mentor_name: Optional[str] = None,
    batch_name: Optional[str] = None,
    course_name: Optional[str] = None,
    session_type: Optional[str] = None,
    recording_status: Optional[str] = None,
    report_status: Optional[str] = None,
    search: Optional[str] = None,
    db: Session = Depends(get_db),
):
    filters = _session_report_filters(
        date_from, date_to, status, mentor_name, batch_name,
        course_name, session_type, recording_status, report_status,
    )
    csv_content = session_reports.build_csv(db, filters, search)

    return Response(
        content=csv_content,
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=session_reports.csv"},
    )


def _session_reports_filter_description(date_from, date_to, status, mentor_name, batch_name,
                                          course_name, session_type, recording_status, report_status, search):
    parts = []
    if date_from:
        parts.append(f"From {date_from}")
    if date_to:
        parts.append(f"To {date_to}")
    if status:
        parts.append(f"Status = {status}")
    if mentor_name:
        parts.append(f"Mentor = {mentor_name}")
    if batch_name:
        parts.append(f"Batch = {batch_name}")
    if course_name:
        parts.append(f"Course = {course_name}")
    if session_type:
        parts.append(f"Type = {session_type}")
    if recording_status:
        parts.append(f"Recording = {recording_status}")
    if report_status:
        parts.append(f"Report = {report_status}")
    if search:
        parts.append(f'Search = "{search}"')
    return "; ".join(parts) if parts else "All data (no filters applied)"


@app.get("/session-reports/export/pdf")
def export_session_reports_pdf(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    status: Optional[str] = None,
    mentor_name: Optional[str] = None,
    batch_name: Optional[str] = None,
    course_name: Optional[str] = None,
    session_type: Optional[str] = None,
    recording_status: Optional[str] = None,
    report_status: Optional[str] = None,
    search: Optional[str] = None,
    db: Session = Depends(get_db),
):
    filters = _session_report_filters(
        date_from, date_to, status, mentor_name, batch_name,
        course_name, session_type, recording_status, report_status,
    )
    rows = session_reports.list_session_reports(db, filters, page=1, page_size=100000, search=search)["items"]
    summary = session_reports.session_reports_summary(db, filters)
    filter_desc = _session_reports_filter_description(
        date_from, date_to, status, mentor_name, batch_name,
        course_name, session_type, recording_status, report_status, search,
    )
    pdf_path = generate_session_reports_list_pdf(rows, summary, filter_desc)

    return FileResponse(path=pdf_path, filename="Session_Reports.pdf", media_type="application/pdf")


@app.get("/session-reports/{session_id}")
def get_session_report_detail(session_id: int, db: Session = Depends(get_db)):
    bundle = session_reports.session_detail_bundle(db, session_id)
    if not bundle:
        return {"success": False, "message": "Session not found"}
    return {"success": True, **bundle}


@app.get("/session-reports/{session_id}/live-details")
def get_session_live_details(session_id: int, db: Session = Depends(get_db)):
    details = session_reports.live_details_bundle(db, session_id)
    if not details:
        return {"success": False, "message": "Session not found"}
    return {"success": True, **details}


@app.get("/session-reports/{session_id}/attendance")
def get_session_attendance(session_id: int, db: Session = Depends(get_db)):
    bundle = session_reports.attendance_bundle(db, session_id)
    if not bundle:
        return {"success": False, "message": "Session not found"}
    return {"success": True, **bundle}


@app.post("/session-reports/{session_id}/attendance")
def create_session_attendance_row(session_id: int, data: SessionAttendanceCreate, db: Session = Depends(get_db)):
    row = session_reports.add_attendance_row(db, session_id, data.dict())
    return {"success": True, "id": row.id}


@app.put("/session-reports/{session_id}/attendance/{attendance_id}")
def update_session_attendance_row(session_id: int, attendance_id: int, data: SessionAttendanceUpdate, db: Session = Depends(get_db)):
    row = session_reports.update_attendance_row(db, attendance_id, data.dict(exclude_unset=True))
    if not row:
        return {"success": False, "message": "Attendance record not found"}
    return {"success": True}


@app.delete("/session-reports/{session_id}/attendance/{attendance_id}")
def delete_session_attendance_row(session_id: int, attendance_id: int, db: Session = Depends(get_db)):
    deleted = session_reports.delete_attendance_row(db, attendance_id)
    if not deleted:
        return {"success": False, "message": "Attendance record not found"}
    return {"success": True}


@app.get("/session-reports/{session_id}/feedback")
def get_session_feedback(session_id: int, db: Session = Depends(get_db)):
    bundle = session_reports.feedback_bundle(db, session_id)
    if not bundle:
        return {"success": False, "message": "Session not found"}
    return {"success": True, **bundle}


@app.post("/session-reports/{session_id}")
def create_session_report(session_id: int, data: SessionReportUpdate, db: Session = Depends(get_db)):
    if not session_reports.get_session(db, session_id):
        return {"success": False, "message": "Session not found"}
    report = session_reports.upsert_report(db, session_id, data.dict(exclude_unset=True))
    return {"success": True, "id": report.id}


@app.put("/session-reports/{session_id}")
def update_session_report(session_id: int, data: SessionReportUpdate, db: Session = Depends(get_db)):
    if not session_reports.get_session(db, session_id):
        return {"success": False, "message": "Session not found"}
    session_reports.upsert_report(db, session_id, data.dict(exclude_unset=True))
    return {"success": True, "message": "Report updated"}


@app.put("/session-reports/{session_id}/status")
def update_session_report_status(session_id: int, data: ReportStatusUpdate, db: Session = Depends(get_db)):
    if not session_reports.get_session(db, session_id):
        return {"success": False, "message": "Session not found"}
    session_reports.update_report_status(db, session_id, data.report_status, data.reviewed_by)
    return {"success": True, "message": "Report status updated"}


@app.get("/session-reports/{session_id}/download")
def download_session_report(session_id: int, db: Session = Depends(get_db)):
    bundle = session_reports.session_detail_bundle(db, session_id)
    if not bundle:
        return {"success": False, "message": "Session not found"}

    attendance = session_reports.attendance_bundle(db, session_id)
    feedback = session_reports.feedback_bundle(db, session_id)

    pdf_path = generate_session_report_pdf(bundle, attendance, feedback)

    return FileResponse(
        path=pdf_path,
        filename=f"Session_Report_{session_id}.pdf",
        media_type="application/pdf",
    )


@app.get("/dashboard/session-reports-today")
def get_dashboard_session_reports_today(db: Session = Depends(get_db)):
    return session_reports.dashboard_today_summary(db)
